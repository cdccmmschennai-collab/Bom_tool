"""Parts Master search over conversion lines + material_master_import.

Runs against PostgreSQL (TEST_DATABASE_URL, wiped by these tests).
"""
import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, func, insert, select
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from app.database import Base, engine

try:
    with engine.connect():
        pass
except OperationalError:  # pragma: no cover
    pytest.skip("PostgreSQL test database not reachable", allow_module_level=True)

from app.main import app  # noqa: E402
from app.manage import main as manage  # noqa: E402
from app.models import ExtractedLine, material_master_import as master  # noqa: E402
from app.services.parts import backfill, like_pattern  # noqa: E402

PASSWORD = "correct-horse-1"


@pytest.fixture(scope="module")
def client():
    Base.metadata.drop_all(engine)
    with TestClient(app) as c:
        mp = pytest.MonkeyPatch()
        mp.setattr("sys.stdin", io.StringIO(PASSWORD + "\n"))
        manage(["create-user", "parts", "--password-stdin"])
        mp.undo()
        assert c.post("/api/auth/login", json={"username": "parts", "password": PASSWORD}).status_code == 200
        yield c


@pytest.fixture(scope="module")
def job(client, sample_bytes):
    dk = client.get("/api/plants", params={"plant_type": "DUKHAN"}).json()[0]
    r = client.post("/api/jobs", data={"plant_type": "DUKHAN", "plant_id": str(dk["id"])},
                    files={"file": ("in.xlsx", sample_bytes, "application/octet-stream")})
    assert r.status_code == 201, r.text
    with engine.begin() as conn:
        conn.execute(insert(master).values(material_temp_number=777001, maintenance_planning_plant="9999",
                                           material_is_equipment_indicator=False, material_type_category="L",
                                           manufacturers_part_number="LEGACY-PN_1", manufacturer_name="ACME"))
    return r.json()


def find(client, field, q, **params):
    r = client.get("/api/parts", params={"field": field, "q": q, **params})
    assert r.status_code == 200, r.text
    return r.json()


def test_like_pattern():
    assert like_pattern("12*") == "12%"
    assert like_pattern("*a_b%*") == "%a\\_b\\%%"


def test_a_part_number_in_both_sources_is_one_row(client, job):
    res = find(client, "material_temp_number", "*", limit=200)
    key = lambda i: " ".join((i["part_number"] or "").split()).upper()  # noqa: E731
    conv = {key(i) for i in res["items"] if i["source"] == "CONVERSION" and i["part_number"]}
    mast = {key(i) for i in res["items"] if i["source"] == "MASTER" and i["part_number"]}
    assert not conv & mast  # never two rows (Conversion + Material master) for one part number
    assert any(i["source"] == "BOTH" for i in res["items"])     # the conversion saved them into the master
    assert any(i["source"] == "MASTER" for i in res["items"])   # master-only rows stay as they are
    merged = [key(i) for i in res["items"] if i["source"] == "BOTH"]
    assert len(merged) == len(set(merged))


def test_merge_fills_columns_from_both_records(client, job):
    with Session(engine, expire_on_commit=False) as db:
        line = db.execute(select(ExtractedLine).where(
            ExtractedLine.job_id == job["id"], ExtractedLine.mfr_part_number.is_not(None),
            ExtractedLine.tag.is_not(None))).scalars().first()
        db.execute(insert(master).values(material_temp_number=888001, maintenance_planning_plant="8888",
                                         material_is_equipment_indicator=False,
                                         manufacturers_part_number=line.mfr_part_number.lower(),
                                         sap_material_number="SAP-MERGE-TEST"))
        db.commit()
    rows = [i for i in find(client, "tag_number", line.tag, limit=200)["items"]
            if (i["part_number"] or "").upper().startswith(line.mfr_part_number.upper())]
    assert len(rows) == 1
    row = rows[0]
    assert row["source"] == "BOTH"
    assert row["tag_number"] == line.tag                 # only the conversion has it
    assert "SAP-MERGE-TEST" in row["sap_material_number"]  # only the material master has it
    assert row["country_name"] == line.country_name     # only the conversion has it
    assert row["planning_plant"].count(line.planning_plant) == 1  # same value shown once
    # the same row when searching from the master side
    by_sap = find(client, "sap_material_number", "SAP-MERGE-TEST")["items"]
    assert [i["source"] for i in by_sap] == ["BOTH"] and line.tag in by_sap[0]["tag_number"]


def test_merge_rows_rules():
    from app.services.parts import merge_rows
    base = dict(job_id=None, created_at=None, material_temp_number=None, material_category=None, description=None,
                planning_plant=None, manufacturer_name=None, spir=None)
    conv = dict(base, source="CONVERSION", job_id=7, part_number="ZB5AV043", tag_number="6130-MC-01",
                sap_material_number=None, country_name="QATAR", planning_plant="2300")
    mast = dict(base, source="MASTER", part_number="ZB5AV043", tag_number=None, sap_material_number="10231210",
                country_name=None, planning_plant="2300")
    row = merge_rows([mast, conv])
    assert (row["source"], row["part_number"], row["tag_number"], row["sap_material_number"],
            row["country_name"], row["planning_plant"], row["job_id"]) == (
        "BOTH", "ZB5AV043", "6130-MC-01", "10231210", "QATAR", "2300", 7)


def test_master_rows_and_literal_underscore(client, job):
    res = find(client, "part_number", "legacy-pn_1")
    assert [(i["source"], i["material_temp_number"]) for i in res["items"]] == [("MASTER", "777001")]
    assert find(client, "part_number", "LEGACY-PNX1")["total"] == 0  # "_" is not a wildcard
    assert find(client, "manufacturer_name", "acme")["total"] == 1


def test_validation_and_paging(client, job):
    assert client.get("/api/parts", params={"field": "nope", "q": "1"}).status_code == 422
    assert client.get("/api/parts", params={"field": "spir", "q": "  "}).status_code == 422
    page = find(client, "material_temp_number", "*", limit=5, offset=5)
    assert len(page["items"]) == 5 and page["total"] > 10


def _export(client, **params):
    import openpyxl
    r = client.get("/api/parts/export", params=params)
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"].startswith("attachment; filename*=UTF-8''PARTS_MASTER_")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert [ws.title.strip() for ws in wb.worksheets] == ["MATERIAL MASTER"]
    ws = wb.worksheets[0]
    data = [row for row in ws.iter_rows(min_row=6, values_only=True) if any(v is not None for v in row)]
    return ws, data


def test_export_search_results(client, job):
    import openpyxl
    from app.config import get_settings
    template = openpyxl.load_workbook(get_settings().submission_dukhan)
    tmpl_ws = next(ws for ws in template.worksheets if ws.title.strip() == "MATERIAL MASTER")

    found = find(client, "tag_number", "34-QD-23", limit=200)
    ws, data = _export(client, field="tag_number", q="34-QD-23")
    for r in range(1, 6):  # same headings / field codes as the submission template
        assert [c.value for c in ws[r]] == [c.value for c in tmpl_ws[r]]
    assert len(data) == found["total"]
    hdr = {str(c.value).strip().upper(): c.column - 1 for c in ws[1] if c.value}
    temps = [row[hdr["MATERIAL TEMP NUMBER"]] for row in data]
    assert all(isinstance(t, int) for t in temps if t is not None and "," not in str(t))
    codes = {row[hdr["MANUFACTURER COUNTRY CODE"]] for row in data} - {None}
    assert codes and all(len(c) <= 3 or "," in c for c in codes)  # names were turned into country codes


def test_export_without_search_is_everything(client, job):
    from app.services.parts import find_parts
    with Session(engine) as db:
        expected = len(find_parts(db))
    _, data = _export(client)
    assert len(data) == expected > 0


# runs last: it deletes the module's conversion
def test_backfill_and_delete(client, job):
    with Session(engine) as db:
        before = db.scalar(select(func.count()).where(ExtractedLine.job_id == job["id"]))
        db.execute(delete(ExtractedLine).where(ExtractedLine.job_id == job["id"]))
        db.commit()
        assert backfill(db) == 1  # read back from the stored working template
        assert db.scalar(select(func.count()).where(ExtractedLine.job_id == job["id"])) == before
    assert find(client, "tag_number", "34-QD-23")["total"] > 0  # tags restored from the tag sheet

    assert client.delete(f"/api/jobs/{job['id']}").status_code == 204
    with Session(engine) as db:
        assert db.scalar(select(func.count()).where(ExtractedLine.job_id == job["id"])) == 0
