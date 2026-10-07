"""History "Combine": stack stored working / submission templates into one workbook.

Runs against PostgreSQL (TEST_DATABASE_URL, wiped by these tests).
"""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.exc import OperationalError

from app.database import Base, engine

try:
    with engine.connect():
        pass
except OperationalError:  # pragma: no cover
    pytest.skip("PostgreSQL test database not reachable", allow_module_level=True)

from app.main import app  # noqa: E402
from app.manage import main as manage  # noqa: E402

PASSWORD = "correct-horse-1"


@pytest.fixture(scope="module")
def client():
    Base.metadata.drop_all(engine)
    with TestClient(app) as c:
        mp = pytest.MonkeyPatch()
        mp.setattr("sys.stdin", io.StringIO(PASSWORD + "\n"))
        manage(["create-user", "combiner", "--password-stdin"])
        mp.undo()
        assert c.post("/api/auth/login", json={"username": "combiner", "password": PASSWORD}).status_code == 200
        yield c


def plant(client, plant_type):
    return client.get("/api/plants", params={"plant_type": plant_type}).json()[0]


def convert(client, content, plant_type):
    r = client.post("/api/jobs", data={"plant_type": plant_type, "plant_id": str(plant(client, plant_type)["id"])},
                    files={"file": ("in.xlsx", content, "application/octet-stream")})
    assert r.status_code == 201, r.text
    return r.json()


def renamed_tags(sample_bytes, suffix):
    wb = openpyxl.load_workbook(io.BytesIO(sample_bytes))
    for row in wb.active.iter_rows(min_row=4):
        if row[2].value:
            row[2].value = f"{row[2].value}-{suffix}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_combine(client, ids, kind):
    r = client.post("/api/jobs/combine", json={"job_ids": ids, "kind": kind})
    assert r.status_code == 202, r.text
    status = client.get(f"/api/jobs/combine/{r.json()['id']}").json()  # the background work has finished
    assert (status["status"], status["done"], status["total"]) == ("DONE", len(ids), len(ids)), status
    f = client.get(f"/api/jobs/combine/{status['id']}/file")
    assert f.status_code == 200
    return status, openpyxl.load_workbook(io.BytesIO(f.content))


def sheet(wb, name):  # template sheet titles can carry trailing spaces
    return next(ws for ws in wb.worksheets if ws.title.strip() == name)


def rows(ws, first):
    return [r for r in ws.iter_rows(min_row=first, values_only=True) if any(v is not None for v in r)]


@pytest.fixture(scope="module")
def two_jobs(client, sample_bytes):
    return [convert(client, renamed_tags(sample_bytes, s), "DUKHAN") for s in ("A", "B")]


def test_combine_working_template(client, two_jobs):
    status, wb = run_combine(client, [j["id"] for j in reversed(two_jobs)], "WORKING")
    assert status["filename"].startswith("BOM_WORKING_TEMPLATE_COMBINED_DUKHAN_2_FILES_")
    ws = sheet(wb, "BOM WORKING TEMPLATE")
    data = rows(ws, 6)
    assert len(data) == sum(j["line_count"] for j in two_jobs)
    hdr = {c.value: c.column for c in ws[1] if c.value}
    serial = hdr["SERIAL NUMBER"]
    assert [ws.cell(6 + i, serial).value for i in range(len(data))] == list(range(1, len(data) + 1))
    # oldest conversion first, and the second file's LEN() formulas point at their own (new) row
    first_b = 6 + two_jobs[0]["line_count"]
    len_cols = [c for h, c in hdr.items() if str(h).upper().startswith("LEN(")]
    assert len_cols and all(f"{first_b}" in str(ws.cell(first_b, c).value) for c in len_cols)
    tags = rows(sheet(wb, "SPIR TAG vs SUBMT"), 3)
    assert len(tags) == sum(j["equipment_count"] for j in two_jobs)


def test_combine_submission_template(client, two_jobs):
    _, wb = run_combine(client, [j["id"] for j in two_jobs], "SUBMISSION")
    eq = sum(j["equipment_count"] for j in two_jobs)
    sp = sum(j["spare_count"] for j in two_jobs)
    assert len(rows(sheet(wb, "BOM HEADER"), 6)) == eq
    assert len(rows(sheet(wb, "BOM PARTS"), 6)) == sp
    assert len(rows(sheet(wb, "MATERIAL MASTER"), 6)) == eq + sp


def test_combine_validation(client, sample_bytes, two_jobs):
    ids = [j["id"] for j in two_jobs]
    post = lambda body: client.post("/api/jobs/combine", json=body)  # noqa: E731
    assert post({"job_ids": ids[:1], "kind": "WORKING"}).status_code == 422          # one file only
    assert post({"job_ids": [ids[0], ids[0]], "kind": "WORKING"}).status_code == 422  # same file twice
    assert post({"job_ids": ids, "kind": "OTHER"}).status_code == 422               # unknown template
    assert post({"job_ids": [ids[0], 999999], "kind": "WORKING"}).status_code == 404
    other = convert(client, renamed_tags(sample_bytes, "O"), "OTHER")
    r = post({"job_ids": [ids[0], other["id"]], "kind": "WORKING"})
    assert r.status_code == 422 and "plant type" in r.json()["detail"]
    bad = client.post("/api/jobs", data={"plant_type": "DUKHAN"},
                      files={"file": ("bad.xlsx", b"not excel", "application/octet-stream")})
    assert bad.status_code == 422
    failed = client.get("/api/jobs", params={"limit": 1}).json()["items"][0]
    assert failed["status"] == "FAILED"
    assert post({"job_ids": [ids[0], failed["id"]], "kind": "WORKING"}).status_code == 422
    assert client.get("/api/jobs/combine/999999").status_code == 404
