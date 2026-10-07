"""End-to-end API tests against PostgreSQL (TEST_DATABASE_URL, wiped by these tests)."""
import io
import re
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.exc import OperationalError

from app.database import Base, engine

try:
    with engine.connect():
        pass
except OperationalError:  # pragma: no cover
    pytest.skip("PostgreSQL test database not reachable", allow_module_level=True)

from app.main import app  # noqa: E402
from app.manage import main as manage  # noqa: E402

PASSWORD = "correct-horse-1"


def make_user(monkeypatch, username, email=None, name=None):
    monkeypatch.setattr("sys.stdin", io.StringIO(PASSWORD + "\n"))
    args = ["create-user", username, "--password-stdin"]
    if email:
        args += ["--email", email]
    if name:
        args += ["--name", name]
    manage(args)


@pytest.fixture(scope="module")
def client():
    Base.metadata.drop_all(engine)
    with TestClient(app) as c:  # runs start-up: create tables, seed plants + countries
        mp = pytest.MonkeyPatch()
        make_user(mp, "tester", email="tester@example.com", name="Test User")
        mp.undo()
        assert c.post("/api/auth/login", json={"username": "tester", "password": PASSWORD}).status_code == 200
        yield c


def plants(client, plant_type):
    return client.get("/api/plants", params={"plant_type": plant_type}).json()


def upload(client, sample_bytes, plant_type, plant_id=None, name="sample.xlsx"):
    data = {"plant_type": plant_type}
    if plant_id:
        data["plant_id"] = str(plant_id)
    return client.post("/api/jobs", data=data, files={"file": (name, sample_bytes, "application/octet-stream")})


def output_sheet(client, job_id):
    r = client.get(f"/api/jobs/{job_id}/output")
    assert r.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(r.content))["BOM WORKING TEMPLATE"]


def test_health_and_reference(client):
    assert client.get("/api/health").json() == {"status": "ok"}
    assert [p["value"] for p in client.get("/api/plant-types").json()] == ["DUKHAN", "OTHER"]
    assert len(plants(client, "DUKHAN")) >= 1 and len(plants(client, "OTHER")) >= 1


def test_convert_and_numbers_continue(client, sample_bytes):
    dk = plants(client, "DUKHAN")[0]
    r = upload(client, sample_bytes, "DUKHAN", dk["id"])
    assert r.status_code == 201, r.text
    job = r.json()
    assert (job["status"], job["equipment_count"], job["spare_count"], job["line_count"]) == ("SUCCESS", 4, 70, 74)
    assert job["spir_numbers"] == ["VEN-4391-M3TY-2-43-0003"]
    assert job["output_filename"].startswith("BOM_WORKING_TEMPLATE_DUKHAN_")
    ws = output_sheet(client, job["id"])
    assert ws["C6"].value == 40001 and ws["C7"].value == 500001
    assert ws["K6"].value == dk["code"] and ws["Y7"].value == "QA"

    # same file again -> same items keep the same numbers
    job2 = upload(client, sample_bytes, "DUKHAN", dk["id"]).json()
    ws2 = output_sheet(client, job2["id"])
    assert ws2["C6"].value == 40001 and ws2["C7"].value == 500001

    # a NEW tag in the same plant continues the sequence (40005, not 40001)
    wb = openpyxl.load_workbook(io.BytesIO(sample_bytes))
    for row in wb.active.iter_rows(min_row=4):
        if row[2].value == "34-QD-23":
            row[2].value = "NEW-TAG-01"
        if row[23].value == "10231307":  # a new part number -> a new spare number
            row[14].value = "NEW-PART-01"
    buf = io.BytesIO()
    wb.save(buf)
    ws3 = output_sheet(client, upload(client, buf.getvalue(), "DUKHAN", dk["id"]).json()["id"])
    assert ws3["C6"].value == 40005 and ws3["C7"].value == 500063

    # another plant continues the same numbering - a temp number is never issued twice
    ot = plants(client, "OTHER")[0]
    job4 = upload(client, sample_bytes, "OTHER", ot["id"]).json()
    ws4 = output_sheet(client, job4["id"])
    assert ws4["B6"].value == 40006 and ws4["B7"].value == 10231307


def test_concurrent_uploads_never_share_numbers(client, sample_bytes):
    ot = plants(client, "OTHER")[-1]

    def run(i):
        wb = openpyxl.load_workbook(io.BytesIO(sample_bytes))
        for row in wb.active.iter_rows(min_row=4):
            if row[2].value:
                row[2].value = f"{row[2].value}-C{i}"
        buf = io.BytesIO()
        wb.save(buf)
        return upload(client, buf.getvalue(), "OTHER", ot["id"]).json()["id"]

    with ThreadPoolExecutor(4) as ex:
        ids = list(ex.map(run, range(4)))
    numbers = []
    for job_id in ids:
        ws = output_sheet(client, job_id)
        numbers += [ws.cell(r, 2).value for r in range(6, 80) if ws.cell(r, 3).value == "B"]
    assert len(numbers) == 16 and len(set(numbers)) == 16


def test_validation_errors(client, sample_bytes):
    ot = plants(client, "OTHER")[0]
    assert upload(client, sample_bytes, "DUKHAN", ot["id"]).status_code == 422  # plant of wrong type
    assert upload(client, sample_bytes, "XYZ").status_code == 422
    assert upload(client, b"not excel", "OTHER", name="bad.xlsx").status_code == 422
    assert upload(client, sample_bytes, "OTHER", name="bad.csv").status_code == 422
    hist = client.get("/api/jobs", params={"limit": 100}).json()
    assert any(j["status"] == "FAILED" for j in hist["items"])


def test_no_plant_selected_leaves_column_empty(client, sample_bytes):
    job = upload(client, sample_bytes, "OTHER").json()
    ws = output_sheet(client, job["id"])
    assert ws["I6"].value is None and ws["I7"].value is None
    assert client.get(f"/api/jobs/{job['id']}/input").status_code == 200


def test_job_records_signed_in_user(client, sample_bytes):
    assert upload(client, sample_bytes, "OTHER").json()["user_name"] == "Test User"


def test_api_requires_sign_in(client):
    with TestClient(app) as anon:
        assert anon.get("/api/health").status_code == 200
        for path in ("/api/plants", "/api/plant-types", "/api/jobs", "/api/jobs/1/output", "/api/auth/me"):
            assert anon.get(path).status_code == 401, path
        anon.cookies.set("bom_session", "forged-token")
        assert anon.get("/api/jobs").status_code == 401


def test_login_logout_and_remember(client, monkeypatch):
    make_user(monkeypatch, "Alice", email="Alice@Example.com")
    with TestClient(app) as c:
        assert c.post("/api/auth/login", json={"username": "alice", "password": "wrong"}).status_code == 401
        assert c.post("/api/auth/login", json={"username": "nobody", "password": PASSWORD}).status_code == 401

        r = c.post("/api/auth/login", json={"username": "ALICE@example.com", "password": PASSWORD})  # e-mail works
        assert r.status_code == 200 and r.json()["username"] == "alice"
        cookie = r.headers["set-cookie"].lower()
        assert "httponly" in cookie and "max-age" not in cookie  # browser-session cookie
        assert c.get("/api/auth/me").json()["username"] == "alice"

        assert c.post("/api/auth/logout").status_code == 204
        assert c.get("/api/auth/me").status_code == 401

        r = c.post("/api/auth/login", json={"username": "alice", "password": PASSWORD, "remember": True})
        assert "max-age=2592000" in r.headers["set-cookie"].lower()  # 30 days

        manage(["disable-user", "alice"])  # disabling signs the user out at once
        assert c.get("/api/auth/me").status_code == 401
        assert c.post("/api/auth/login", json={"username": "alice", "password": PASSWORD}).status_code == 401


def test_profile(client):
    p = client.get("/api/auth/profile").json()
    assert p["username"] == "tester" and p["full_name"] == "Test User" and p["email"] == "tester@example.com"
    jobs = client.get("/api/jobs", params={"limit": 100}).json()["items"]
    mine = [j for j in jobs if j["user_name"] == "Test User" and j["status"] == "SUCCESS"]
    assert p["extraction_count"] == len(mine)
    assert p["last_login"]


def test_change_password(client, monkeypatch):
    make_user(monkeypatch, "bob", name="Bob Smith")
    new = "brand-new-pass-2"
    with TestClient(app) as here, TestClient(app) as other:
        for c in (here, other):
            assert c.post("/api/auth/login", json={"username": "bob", "password": PASSWORD}).status_code == 200
        assert here.get("/api/auth/profile").json()["extraction_count"] == 0

        def change(current, new_password):
            return here.post("/api/auth/change-password",
                             json={"current_password": current, "new_password": new_password})

        assert change("wrong-password", new).status_code == 400
        assert change(PASSWORD, "short").status_code == 400
        assert change(PASSWORD, PASSWORD).status_code == 400
        assert here.get("/api/auth/me").status_code == 200 and other.get("/api/auth/me").status_code == 200

        assert change(PASSWORD, new).status_code == 204
        assert here.get("/api/auth/me").status_code == 200   # this device stays signed in
        assert other.get("/api/auth/me").status_code == 401  # other devices are signed out
        assert other.post("/api/auth/login", json={"username": "bob", "password": PASSWORD}).status_code == 401
        assert other.post("/api/auth/login", json={"username": "bob", "password": new}).status_code == 200

    with TestClient(app) as anon:
        assert anon.get("/api/auth/profile").status_code == 401
        assert anon.post("/api/auth/change-password",
                         json={"current_password": PASSWORD, "new_password": new}).status_code == 401


def test_history_search_and_dates(client, sample_bytes):
    jobs = client.get("/api/jobs", params={"limit": 200}).json()
    assert jobs["total"] > 0
    spir = next(j for j in jobs["items"] if j["spir_numbers"])["spir_numbers"][0]
    found = client.get("/api/jobs", params={"limit": 200, "search": spir}).json()
    assert found["total"] > 0 and all(spir in (j["spir_numbers"] or []) for j in found["items"])

    newest = jobs["items"][0]["created_at"]
    day = newest[:10]
    assert client.get("/api/jobs", params={"date_from": "2000-01-01T00:00:00Z"}).json()["total"] == jobs["total"]
    assert client.get("/api/jobs", params={"date_to": "2000-01-01T00:00:00Z"}).json()["total"] == 0
    assert client.get("/api/jobs", params={"date_from": "2999-01-01T00:00:00Z"}).json()["total"] == 0
    same_day = client.get("/api/jobs", params={"date_from": f"{day}T00:00:00Z", "date_to": f"{day}T23:59:59.999Z",
                                               "limit": 200}).json()
    assert same_day["total"] > 0 and all(j["created_at"][:10] == day for j in same_day["items"])
    assert client.get("/api/jobs", params={"date_from": "not-a-date"}).status_code == 422


def submission_book(client, job_id):
    r = client.get(f"/api/jobs/{job_id}/submission")
    assert r.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(r.content))


def data_rows(ws, first=6):
    return [[c.value for c in row] for row in ws.iter_rows(min_row=first) if any(c.value is not None for c in row)]


@pytest.mark.parametrize("plant_type", ["DUKHAN", "OTHER"])
def test_submission_template(client, sample_bytes, plant_type):
    pl = plants(client, plant_type)[0]
    job = upload(client, sample_bytes, plant_type, pl["id"]).json()
    assert job["submission_filename"].startswith(
        f"BOM_SUBMISSION_TEMPLATE_{'DUKHAN' if plant_type == 'DUKHAN' else 'OTHER_PLANT'}_")
    wb = submission_book(client, job["id"])
    assert wb.sheetnames == ["BOM HEADER", "BOM PARTS", "MATERIAL MASTER"]
    work = output_sheet(client, job["id"])
    wcol = {re.sub(r"[^A-Z0-9]", "", str(c.value).upper()): c.column - 1 for c in work[1] if c.value}
    col = lambda name: wcol[re.sub(r"[^A-Z0-9]", "", name.upper())]  # noqa: E731
    working = data_rows(work)

    header, parts, master = (data_rows(wb[n]) for n in wb.sheetnames)
    assert len(header) == job["equipment_count"] and len(parts) == job["spare_count"]
    assert len(master) == job["line_count"]

    # BOM HEADER: equipment temp numbers, equipment indicator YES, empty remark columns
    equipment = [w for w in working if w[col("MATERIAL IS AN EQUIPMENT INDICATOR")] == "YES"]
    assert [h[0] for h in header] == [w[col("MATERIAL TEMP NUMBER")] for w in equipment]
    assert all(h[1] == "B" and h[11] == "YES" and h[5] == pl["code"] for h in header)
    assert all(h[13] is None and h[14] is None for h in header)
    # descriptions line up (the Dukhan template ships a stray value in E6 that must be overwritten/cleared)
    assert [h[4] for h in header] == [w[col("MATERIAL DESCRIPTION")] for w in equipment]

    # BOM PARTS: every spare points at an equipment in BOM HEADER
    header_numbers = {h[0] for h in header}
    assert all(p[0] in header_numbers and p[2] == "L" for p in parts)
    spares = [w for w in working if w[col("MATERIAL IS AN EQUIPMENT INDICATOR")] == "NO"]
    assert [(p[1], p[12], p[13]) for p in parts] == [
        (w[col("MATERIAL TEMP NUMBER")], w[col("QUANTITY")], w[col("POSITION NUMBER")]) for w in spares]
    assert all(p[15] is None for p in parts)  # CDC_BOM Parts Remark left empty

    # MATERIAL MASTER: every line, in working-template order
    assert [m[0] for m in master] == [w[col("MATERIAL TEMP NUMBER")] for w in working]
    assert wb["MATERIAL MASTER"].freeze_panes == "A6"


def test_submission_missing_for_old_or_failed_jobs(client, sample_bytes):
    failed = upload(client, b"not excel", "OTHER", name="bad.xlsx")
    assert failed.status_code == 422
    hist = client.get("/api/jobs", params={"limit": 1}).json()["items"][0]
    assert hist["status"] == "FAILED" and hist["submission_filename"] is None
    assert client.get(f"/api/jobs/{hist['id']}/submission").status_code == 404
    assert client.get("/api/jobs/999999/submission").status_code == 404


def test_delete_job_keeps_issued_numbers(client, sample_bytes):
    dk = plants(client, "DUKHAN")[0]
    wb = openpyxl.load_workbook(io.BytesIO(sample_bytes))
    for row in wb.active.iter_rows(min_row=4):
        if row[2].value:
            row[2].value = f"{row[2].value}-DEL"  # new tags -> this job issues new equipment numbers
    buf = io.BytesIO()
    wb.save(buf)
    job = upload(client, buf.getvalue(), "DUKHAN", dk["id"]).json()
    first = output_sheet(client, job["id"])["C6"].value

    assert client.delete(f"/api/jobs/{job['id']}").status_code == 204
    assert client.get(f"/api/jobs/{job['id']}").status_code == 404
    assert client.delete(f"/api/jobs/{job['id']}").status_code == 404

    # the numbers stay taken: the same tags get the same numbers again
    again = upload(client, buf.getvalue(), "DUKHAN", dk["id"]).json()
    assert output_sheet(client, again["id"])["C6"].value == first
