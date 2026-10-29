"""Business-logic tests (no database needed)."""
import io

import openpyxl
import pytest

from app.config import get_settings
from app.services.bom_builder import (DUKHAN, OTHER, InMemoryAllocator, build_bom, code_part,
                                      split_min_max, split_spir)
from app.services.excel_writer import write_workbook
from app.services.input_reader import InputFileError, read_input

COUNTRIES = {"QATAR": "QA"}


def build(sample_bytes, plant_type, plant="DK01"):
    return build_bom(read_input(sample_bytes), plant_type, plant, InMemoryAllocator(), COUNTRIES)


def test_helpers():
    assert split_spir("VEN-4391-M3TY-2-43-0003-REV-A") == ("VEN-4391-M3TY-2-43-0003", "A")
    assert split_spir("VEN-4391-M3TY-2-43-0003") == ("VEN-4391-M3TY-2-43-0003", None)
    assert code_part("NO - NUMBERS") == "NO"
    assert code_part("QAR - QATAR RIYAL") == "QAR"
    assert split_min_max("1") == (1, None)
    assert split_min_max("0") == (0, None)
    assert split_min_max("12") == (None, 12)
    assert split_min_max("120") == (None, 120)
    assert split_min_max("1/3") == (1, 3)


def test_reader_skips_meta_rows(sample_bytes):
    rows = read_input(sample_bytes)
    assert len(rows) == 78
    assert rows[0].get("TAG NO") == "34-QD-23"


def test_reader_rejects_wrong_file():
    wb = openpyxl.Workbook()
    wb.active.append(["A", "B"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(InputFileError, match="Missing columns"):
        read_input(buf.getvalue())


def test_dukhan_lines(sample_bytes):
    res = build(sample_bytes, DUKHAN)
    assert (res.equipment_count, res.spare_count, len(res.lines)) == (4, 70, 74)
    eq = res.lines[0]
    assert (eq.category, eq.equipment_indicator, eq.temp_number, eq.bom_header_number) == ("B", "YES", 40001, 40001)
    assert eq.description == "LV SWITCHGEAR" and eq.quantity == 1 and eq.planning_plant == "DK01"
    sp = res.lines[1]
    assert sp.category == "L" and sp.equipment_indicator == "NO"
    assert sp.sap_material_number == 10231307 and sp.temp_number == 10231307  # SAP number used in Dukhan too
    assert sp.bom_header_number == 40001 and sp.quantity == 4 and sp.position_number == "0010"
    assert sp.spir_number == "VEN-4391-M3TY-2-43-0003"
    assert sp.uom == "NO" and sp.currency == "QAR" and sp.country_code == "QA" and sp.country_name == "QATAR"
    # every equipment line is EA, whatever its input UNIT OF MEASURE; spares keep their own
    assert {line.uom for line in res.lines if line.category == "B"} == {"EA"}
    assert sp.price == 75.3 and sp.delivery_weeks == 20 and sp.min_qty == 1 and sp.max_qty is None
    # same material under several tags keeps one temp number
    by_sap = {}
    for line in res.lines:
        if line.category == "L":
            by_sap.setdefault(line.sap_material_number, set()).add(line.temp_number)
    assert all(len(v) == 1 for v in by_sap.values())
    assert len(by_sap) == 62
    # tags repeated over several input sheets are merged into one equipment
    assert [t.tag for t in res.tags] == ["34-QD-23", "6130-MC-01", "34-MCC-001", "1660-MC-01"]
    assert sum(1 for line in res.lines if line.bom_header_number == 40002) == 21


def test_other_plant_uses_sap_number(sample_bytes):
    res = build(sample_bytes, OTHER, None)
    assert res.lines[0].temp_number == 40001
    assert res.lines[1].temp_number == 10231307
    assert res.lines[1].planning_plant is None
    assert res.lines[1].sap_material_number == 10231307  # kept on the line, but no column in this template


def test_tag_sheet(sample_bytes):
    t = build(sample_bytes, DUKHAN).tags[0]
    assert (t.tag, t.submt, t.spir_no, t.rev, t.herst, t.typbz, t.mapar, t.serge, t.spir_type) == (
        "34-QD-23", 40001, "VEN-4391-M3TY-2-43-0003", "A", "SCHNEIDER ELECTRIC", "OKKEN", None, "A88029",
        "NORMAL OPERATING SPARES")


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_templates_spir_number_keeps_revision(sample_bytes, plant_type):
    """CDC_SPIR NUMBER on the working and submission templates shows SPIR NO as in the input; the tag sheet,
    the material master and the file names keep the base number."""
    from app.services.excel_writer import TAG_SHEET, _sheet, write_submission

    s = get_settings()
    res = build(sample_bytes, plant_type)
    full, base = "VEN-4391-M3TY-2-43-0003-REV-A", "VEN-4391-M3TY-2-43-0003"
    assert {line.spir_number_full for line in res.lines} == {full}
    assert {line.spir_number for line in res.lines} == {base} and res.spir_numbers == [base]

    def column(ws, header, first_row):
        col = next(c.column for row in ws.iter_rows(max_row=first_row - 1) for c in row
                   if c.value and " ".join(str(c.value).split()).upper() == header)
        return {ws.cell(r, col).value for r in range(first_row, ws.max_row + 1)} - {None}

    wb = openpyxl.load_workbook(io.BytesIO(write_workbook(
        res, plant_type, s.template_dukhan if plant_type == DUKHAN else s.template_other)))
    assert column(wb["BOM WORKING TEMPLATE"], "CDC_SPIR NUMBER", 6) == {full}
    assert column(_sheet(wb, TAG_SHEET), "SPIR NO", 3) == {base}
    bom = wb["BOM WORKING TEMPLATE"]
    hdr = {" ".join(str(c.value).split()).upper(): c.column for c in bom[1] if c.value}
    eq_uom = {bom.cell(r, hdr["BASE UNIT OF MEASURE(T006)"]).value for r in range(6, bom.max_row + 1)
              if bom.cell(r, hdr["MATERIAL IS AN EQUIPMENT INDICATOR"]).value == "YES"}
    assert eq_uom == {"EA"}
    sub = openpyxl.load_workbook(io.BytesIO(write_submission(
        res, s.submission_dukhan if plant_type == DUKHAN else s.submission_other)))
    for name in ("BOM HEADER", "BOM PARTS", "MATERIAL MASTER"):
        assert column(sub[name], "CDC_SPIR NUMBER", 6) == {full}
    assert column(sub["BOM HEADER"], "BASE UNIT OF MEASURE(T006)", 6) == {"EA"}


def _styles(ws, rows):
    out = {}
    for r in rows:
        for c in ws[r]:
            out[c.coordinate] = (c.fill.fill_type, c.fill.fgColor.rgb, c.fill.fgColor.theme, c.fill.fgColor.tint,
                                 c.font.color.rgb if c.font.color else None, c.font.b)
    return out


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_template_colours_unchanged(sample_bytes, plant_type):
    s = get_settings()
    tpl_path = s.template_dukhan if plant_type == DUKHAN else s.template_other
    res = build(sample_bytes, plant_type)
    out = openpyxl.load_workbook(io.BytesIO(write_workbook(res, plant_type, tpl_path)))
    tpl = openpyxl.load_workbook(tpl_path)

    t1, o1 = tpl["BOM WORKING TEMPLATE"], out["BOM WORKING TEMPLATE"]
    assert _styles(t1, range(1, 6)) == _styles(o1, range(1, 6))
    # the template's own conditional formatting is kept as-is; nothing is added
    # (openpyxl renumbers dxfId on save, so compare the actual differential style - fill/font/border)
    cf = lambda ws: sorted((str(k.sqref), [(r.type, r.formula, repr(ws.parent._differential_styles[r.dxfId]))  # noqa: E731
                                           for r in v])
                           for k, v in ws.conditional_formatting._cf_rules.items())
    assert cf(o1) == cf(t1)
    # data cells carry no new fill
    assert all(repr(c.fill) == repr(t1.cell(c.row, c.column).fill)
               for row in o1.iter_rows(min_row=6, max_row=79) for c in row)

    # tag sheet: original columns A-H keep their style, REMARKS moved one column right unchanged
    t2, o2 = tpl.worksheets[1], out.worksheets[1]
    for r in (1, 2):
        for col in range(1, 9):
            assert _styles(t2, [r])[t2.cell(r, col).coordinate] == _styles(o2, [r])[o2.cell(r, col).coordinate]
        assert o2.cell(r, 9).value == "SPIR TYPE"
        assert o2.cell(r, 10).value == t2.cell(r, 9).value  # REMARKS
        assert repr(o2.cell(r, 10).fill) == repr(t2.cell(r, 9).fill)
        assert repr(o2.cell(r, 9).fill) == repr(t2.cell(r, 8).fill)  # new column uses neighbour's header style


def test_written_values(sample_bytes):
    s = get_settings()
    res = build(sample_bytes, DUKHAN)
    ws = openpyxl.load_workbook(io.BytesIO(write_workbook(res, DUKHAN, s.template_dukhan)))["BOM WORKING TEMPLATE"]
    hdr = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    assert ws.cell(7, hdr["SAP MATERIAL NUMBER"]).value == 10231307
    assert ws.cell(7, hdr["MATERIAL TEMP NUMBER"]).value == 10231307  # the SAP number
    assert ws.cell(7, hdr["LEN(MAT.DESC = 40)"]).value == "=LEN(I7)"
    assert ws.cell(79, hdr["SERIAL NUMBER"]).value == 74
    assert ws.cell(80, hdr["SERIAL NUMBER"]).value is None


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_borders_and_frozen_header(sample_bytes, plant_type):
    s = get_settings()
    tpl_path = s.template_dukhan if plant_type == DUKHAN else s.template_other
    res = build(sample_bytes, plant_type)
    out = openpyxl.load_workbook(io.BytesIO(write_workbook(res, plant_type, tpl_path)))
    bom, tags = out["BOM WORKING TEMPLATE"], out.worksheets[1]

    assert bom.freeze_panes == "A6"   # rows 1-5 frozen
    assert tags.freeze_panes == "A3"  # rows 1-2 frozen

    def bordered(c):
        sides = (getattr(c.border, side) for side in ("left", "right", "top", "bottom"))
        return all(s is not None and s.style == "thin" for s in sides)

    for ws, first, last in ((bom, 6, 5 + len(res.lines)), (tags, 3, 2 + len(res.tags))):
        last_col = max(c.column for c in ws[1] if c.value is not None)
        table = [c for row in ws.iter_rows(min_row=first, max_row=last, max_col=last_col) for c in row]
        assert any(c.value is None for c in table)  # the table has empty cells, and they are bordered too
        assert all(bordered(c) for c in table)
        below = [c for row in ws.iter_rows(min_row=last + 1, max_row=ws.max_row) for c in row]
        assert not any(bordered(c) for c in below)  # the grid ends at the last data row


def test_spares_are_numbered_by_part_number():
    from app.services.input_reader import InputRow

    def row(n, tag, pn=None, **extra):
        values = {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": "1", "MANUFACTURER PART NUMBER": pn, **extra}
        return InputRow(n, values)

    rows = [
        InputRow(1, {"TAG NO": "T-1", "EQPT QTY": "1"}),
        row(2, "T-1", "ABC123"),
        row(3, "T-1", "XYZ456"),
        InputRow(4, {"TAG NO": "T-2", "EQPT QTY": "1"}),
        row(5, "T-2", " abc123 "),                           # same part (case / spaces), other tag
        row(6, "T-2", "XYZ 456"),                            # a different part number
        row(7, "T-2", "ABC123", **{"SAP NUMBER": "999"}),    # a SAP NUMBER is used as it is (every plant)
    ]
    spares = [line.temp_number for line in build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {}).lines
              if line.category == "L"]
    assert spares == [500001, 500002, 500001, 500003, 999]


def test_equipment_is_numbered_by_part_number():
    from app.services.input_reader import InputRow

    def eq(n, tag, pn=None):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "MANUFACTURER PART NUMBER": pn})

    rows = [
        eq(1, "T-1", "JB82"),
        eq(2, "T-2", " jb82 "),     # same part (case / spaces), other tag -> same number
        eq(3, "T-3", "P50P9"),      # a different part number
        eq(4, "T-4"),               # no part number -> by tag
        eq(5, "T-5"),
        InputRow(6, {"TAG NO": "T-6", "QUANTITY IDENTICAL PARTS FITTED": "1",  # no equipment row: built from
                     "MANUFACTURER PART NUMBER": "JB82"}),                     # a spare -> by tag, not the spare's PN
    ]
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    equipment = {line.tag: line.temp_number for line in res.lines if line.category == "B"}
    assert equipment == {"T-1": 40001, "T-2": 40001, "T-3": 40002, "T-4": 40003, "T-5": 40004, "T-6": 40005}
    assert [t.submt for t in res.tags] == [40001, 40001, 40002, 40003, 40004, 40005]


def test_equipment_is_numbered_by_make_and_model():
    from app.services.input_reader import InputRow

    def eq(n, tag, make, model, pn=None):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "DESCRIPTION OF PARTS": "CENTRIFUGAL PUMP",
                            "EQPT MAKE": make, "EQPT MODEL": model, "MANUFACTURER PART NUMBER": pn})

    rows = [eq(1, "P-1", "KSB", "ETA 50", pn="A1"),
            eq(2, "P-2", "ksb ", "eta  50", pn="B2"),  # same make + model (case / spaces) -> same number
            eq(3, "P-3", "KSB", "ETA 65"),             # other model -> new number
            eq(4, "P-4", "FLOWSERVE", "ETA 50"),       # other make -> new number
            eq(5, "P-5", "KSB", None, pn="A1")]         # no model -> by part number
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [t.submt for t in res.tags] == [40001, 40001, 40002, 40003, 40004]


def test_equipment_and_spare_keep_separate_numbers_and_reuse_them():
    from app.services.input_reader import InputRow

    def eq(n, tag, pn):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "MANUFACTURER PART NUMBER": pn})

    def sp(n, tag, pn):
        return InputRow(n, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": "1", "MANUFACTURER PART NUMBER": pn})

    rows = [
        eq(1, "T-1", "EBAR10"),
        sp(2, "T-1", "EBAR10"),     # same part number as its equipment -> still a spare number of its own
        sp(3, "T-1", "O-RING"),
        eq(4, "T-2", "EBAR10"),     # the same equipment again (other tag) -> its number again
        sp(5, "T-2", "EBAR10"),     # the same spare again -> its number again
        sp(6, "T-2", "o-ring "),    # the same spare again (case / spaces) -> its number again
        eq(7, "T-3", "P50P9"),      # a new equipment -> new number
        sp(8, "T-3", "GASKET"),     # a new spare -> new number
    ]
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [(line.category, line.temp_number) for line in res.lines] == [
        ("B", 40001), ("L", 500001), ("L", 500002),
        ("B", 40001), ("L", 500001), ("L", 500002),
        ("B", 40002), ("L", 500003)]


def test_manufacturer_name_fits_30_whole_words():
    from app.services.bom_builder import fit_words
    from app.services.input_reader import InputRow

    assert fit_words("EMERSON PROCESS AUTOMATION, QATAR", 30) == "EMERSON PROCESS AUTOMATION"
    assert fit_words("SIEMENS", 30) == "SIEMENS"                      # short values are untouched
    assert fit_words("A" * 30, 30) == "A" * 30                        # exactly 30 is kept
    assert fit_words("ABCDEFGHIJ KLMNOPQRST UVWXYZABC D", 30) == "ABCDEFGHIJ KLMNOPQRST"
    assert fit_words("X" * 35, 30) == "X" * 30                        # one over-long word: length rule wins
    assert fit_words(None, 30) is None

    name = "EMERSON PROCESS AUTOMATION, QATAR"
    rows = [
        InputRow(1, {"TAG NO": "T-1", "EQPT QTY": "1", "EQPT MAKE": name, "SUPPLIER/ OCM NAME": name}),
        InputRow(2, {"TAG NO": "T-1", "QUANTITY IDENTICAL PARTS FITTED": "1", "MANUFACTURER PART NUMBER": "P1",
                     "SUPPLIER/ OCM NAME": name}),
    ]
    result = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [line.mfr_name for line in result.lines] == ["EMERSON PROCESS AUTOMATION"] * 2
    assert result.tags[0].herst == name                               # TAG sheet keeps the full name
    assert [w.lines for w in result.warnings if w.code == "MFR_NAME_TOO_LONG"] == [[1, 2]]
    assert [line.add_information_text for line in result.lines] == [f"Manufacture name : {name}"] * 2


def test_equipment_manufacturer_name_is_eqpt_make():
    from app.services.input_reader import InputRow

    rows = [
        InputRow(1, {"TAG NO": "T-1", "EQPT QTY": "1", "EQPT MAKE": "JOBIRD",
                     "SUPPLIER/ OCM NAME": "INTERNATIONAL GULF TRADING"}),
        InputRow(2, {"TAG NO": "T-1", "QUANTITY IDENTICAL PARTS FITTED": "1", "EQPT MAKE": "JOBIRD",
                     "SUPPLIER/ OCM NAME": "INTERNATIONAL GULF TRADING"}),
        InputRow(3, {"TAG NO": "T-2", "EQPT QTY": "1", "SUPPLIER/ OCM NAME": "SIEMENS"}),  # no EQPT MAKE
        InputRow(4, {"TAG NO": "T-3", "QUANTITY IDENTICAL PARTS FITTED": "1", "EQPT MAKE": "ABB",  # no equipment
                     "SUPPLIER/ OCM NAME": "ABB QATAR"}),                                         # row (no EQPT QTY)
    ]
    result = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [(line.tag, line.category, line.mfr_name) for line in result.lines] == [
        ("T-1", "B", "JOBIRD"), ("T-1", "L", "INTERNATIONAL GULF TRADING"),
        ("T-2", "B", None),
        ("T-3", "B", None), ("T-3", "L", "ABB QATAR")]
    assert [t.herst for t in result.tags] == ["JOBIRD", None, None]  # tag sheet: same EQPT MAKE rule


def test_equipment_manufacturer_name_from_identical_spare():
    from app.services.input_reader import InputRow

    vendor = "INTERNATIONAL GULF TRADING COMPANY /SALAM PETROLEUM"

    def eq(n, tag):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "EQPT MAKE": vendor, "MANUFACTURER PART NUMBER": "JB83BA"})

    def sp(n, tag, pn, name):
        return InputRow(n, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": "1", "MANUFACTURER PART NUMBER": pn,
                            "SUPPLIER/ OCM NAME": name})

    rows = [eq(1, "T-1"), sp(2, "T-1", "jb83ba ", "JOBIRD"),   # the spare is the equipment itself -> its make
            eq(3, "T-2"), sp(4, "T-2", "JB81", "JOBIRD"),      # other part number -> EQPT MAKE
            eq(5, "T-3"), sp(6, "T-3", "JB83BA", None)]        # same part, no SUPPLIER/ OCM NAME -> EQPT MAKE
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    equipment = [ln for ln in res.lines if ln.category == "B"]
    assert [ln.mfr_name for ln in equipment] == ["JOBIRD", "INTERNATIONAL GULF TRADING", "INTERNATIONAL GULF TRADING"]
    assert [ln.add_information_text for ln in equipment] == [None] + [f"Manufacture name : {vendor}"] * 2
    assert [t.herst for t in res.tags] == [vendor] * 3                # tag sheet keeps EQPT MAKE

    # the spare names the same part number after a 'P/N:' label -> still the equipment itself
    rows = [InputRow(1, {"TAG NO": "SCBA-1", "EQPT QTY": "1", "EQPT MAKE": vendor, "MANUFACTURER PART NUMBER": "3357774"}),
            sp(2, "SCBA-1", "DRAGER PAS LITE-P/N: 3357774", "DRÄGER")]
    eq_line, spare_line = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {}).lines
    assert eq_line.mfr_name == "DRÄGER"
    assert spare_line.mfr_part_number == "DRAGER PAS LITE-P/N: 3357774"   # written as in the input


def test_core_part_number_reads_the_value_after_a_label():
    from app.services.bom_builder import _core_part_number

    assert _core_part_number("DRAGER PAS LITE-P/N: 3357774") == "3357774"
    assert _core_part_number("pump pn 12-ab") == "12-AB"
    assert _core_part_number("SEAL PART NO. X1") == "X1"
    assert _core_part_number("PART NUMBER # 77") == "77"
    assert _core_part_number("PN50") == "PN50"                          # not a label
    assert _core_part_number("SPN-12") == "SPN-12"
    assert _core_part_number("1sbl  137001r1110") == "1SBL 137001R1110"
    assert _core_part_number("P/N:") == "P/N:"                          # nothing after the label -> unchanged


def test_add_information_keeps_part_number_and_full_name():
    from app.services.bom_builder import part_number_from_add_information
    from app.services.input_reader import InputRow

    long_pn = "P" * 40
    long_name = "EMERSON PROCESS AUTOMATION, QATAR"
    rows = [
        InputRow(1, {"TAG NO": "T-1", "EQPT QTY": "1", "SUPPLIER/ OCM NAME": "SIEMENS"}),
        InputRow(2, {"TAG NO": "T-1", "QUANTITY IDENTICAL PARTS FITTED": "1", "MANUFACTURER PART NUMBER": long_pn,
                     "SUPPLIER/ OCM NAME": long_name}),
        InputRow(3, {"TAG NO": "T-1", "QUANTITY IDENTICAL PARTS FITTED": "1", "MANUFACTURER PART NUMBER": long_pn,
                     "SUPPLIER/ OCM NAME": "SIEMENS"}),
    ]
    eq, both, pn_only = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {}).lines
    assert eq.add_information_text is None                       # nothing to add -> empty, as before
    assert pn_only.add_information_text == long_pn               # part number only -> unchanged
    assert both.add_information_text == f"{long_pn} | Manufacture name : {long_name}"
    assert both.add_information == long_pn                       # material master still gets the part number
    for line in (eq, both, pn_only):
        assert part_number_from_add_information(line.add_information_text) == line.add_information
    assert part_number_from_add_information(f"Manufacture name : {long_name}") is None


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_one_selected_tab(sample_bytes, plant_type):
    """Two selected tabs open in Excel's [Group] mode, where Filter and Sort are disabled."""
    from app.services.combiner import SUBMISSION, combine
    from app.services.excel_writer import write_submission

    settings = get_settings()
    result = build(sample_bytes, plant_type)
    working = write_workbook(result, plant_type,
                             settings.template_dukhan if plant_type == DUKHAN else settings.template_other)
    submission = write_submission(result,
                                  settings.submission_dukhan if plant_type == DUKHAN else settings.submission_other)

    # a submission made before the fix: BOM HEADER active, MATERIAL MASTER still selected from the template
    old = openpyxl.load_workbook(io.BytesIO(submission))
    for ws in old.worksheets:
        ws.sheet_view.tabSelected = True
    buf = io.BytesIO()
    old.save(buf)

    for content, active in ((working, "BOM WORKING TEMPLATE"), (submission, "BOM HEADER"),
                            (combine([buf.getvalue(), submission], SUBMISSION), "BOM HEADER")):
        wb = openpyxl.load_workbook(io.BytesIO(content))
        assert [ws.title.strip() for ws in wb.worksheets if ws.sheet_view.tabSelected] == [active]
        assert wb.active.title.strip() == active


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_data_text_is_black(sample_bytes, plant_type):
    """The template pre-colours some LEN() cells red; every value in the data table must be black."""
    from app.services.combiner import SUBMISSION, WORKING, combine
    from app.services.excel_writer import write_submission

    settings = get_settings()
    result = build(sample_bytes, plant_type)
    working = write_workbook(result, plant_type,
                             settings.template_dukhan if plant_type == DUKHAN else settings.template_other)
    submission = write_submission(result,
                                  settings.submission_dukhan if plant_type == DUKHAN else settings.submission_other)

    def not_black(content):
        wb = openpyxl.load_workbook(io.BytesIO(content))
        bad = []
        for ws in wb.worksheets:
            first = 3 if ws.title.strip().upper() == "SPIR TAG VS SUBMT" else 6
            for row in ws.iter_rows(min_row=first):
                for c in row:
                    color = c.font.color
                    if c.value is not None and color is not None and not (color.type == "rgb" and color.rgb == "FF000000"):
                        bad.append(f"{ws.title}!{c.coordinate}")
        return bad

    for content in (working, submission, combine([working, working], WORKING),
                    combine([submission, submission], SUBMISSION)):
        assert not_black(content) == []

    # the header block keeps its template colours
    tpl = openpyxl.load_workbook(settings.template_dukhan if plant_type == DUKHAN else settings.template_other)
    out = openpyxl.load_workbook(io.BytesIO(working))
    for row in tpl.worksheets[0].iter_rows(max_row=5):
        for c in row:
            assert out.worksheets[0][c.coordinate].font.color == c.font.color


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_equipment_old_material_is_qequ_plus_temp_number(sample_bytes, plant_type):
    res = build(sample_bytes, plant_type)
    equipment = [l for l in res.lines if l.category == "B"]
    assert equipment
    for line in equipment:
        assert line.old_material_number == f"QEQU-{line.temp_number}"
    # spares keep the OLD MATERIAL NUMBER/SPF NUMBER of their input row
    rows = {r.excel_row: r for r in read_input(sample_bytes)}
    for line in res.lines:
        if line.category == "L":
            assert line.old_material_number == rows[line.input_row].get("OLD MATERIAL NUMBER/SPF NUMBER")


def test_templates_list_each_unique_record_once():
    from app.services.excel_writer import TAG_SHEET, _sheet, write_submission
    from app.services.input_reader import InputRow

    def eq(n, tag):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "MANUFACTURER PART NUMBER": "EBAR10-2-T-316",
                            "DESCRIPTION OF PARTS": "EARTH BAR"})

    def sp(n, tag, pn, qty="1"):
        return InputRow(n, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": qty, "MANUFACTURER PART NUMBER": pn,
                            "DESCRIPTION OF PARTS": f"PART {pn}"})

    rows = []
    for i, tag in enumerate(["24-EB-0068", "24-EB-0069", "24-EB-0072"]):   # the same block three times
        rows += [eq(10 * i + 1, tag), sp(10 * i + 2, tag, "EBAR10-2-T-316"), sp(10 * i + 3, tag, "ELINK-T")]
    rows += [eq(41, "24-EB-0073"), sp(42, "24-EB-0073", "ELINK-T", qty="2"),  # same block, other quantity
             sp(43, "24-EB-0073", "NUT")]                                     # and a spare not listed yet
    result = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert len(result.lines) == 12 and len(result.tags) == 4                  # the conversion keeps every line

    s = get_settings()
    ws = openpyxl.load_workbook(io.BytesIO(write_workbook(result, DUKHAN, s.template_dukhan)))
    bom = ws["BOM WORKING TEMPLATE"]
    hdr = {str(c.value).strip(): c.column for c in bom[1] if c.value}
    got = [(bom.cell(r, hdr["SERIAL NUMBER"]).value, bom.cell(r, hdr["MATERIAL TEMP NUMBER"]).value,
            bom.cell(r, hdr["QUANTITY"]).value) for r in range(6, 11)]
    # one equipment block: B, its two spares (first quantity kept) and the new spare - renumbered 1..4
    assert got == [(1, 40001, 1), (2, 500001, 1), (3, 500002, 1), (4, 500003, 1), (None, None, None)]
    assert [w.lines for w in result.warnings if w.code == "QUANTITY_DIFFERS"] == [[42]]
    tag_sheet = _sheet(ws, TAG_SHEET)
    assert [tag_sheet.cell(r, 1).value for r in range(3, 7)] == ["24-EB-0068", "24-EB-0069", "24-EB-0072",
                                                                 "24-EB-0073"]   # every tag still mapped

    sub = openpyxl.load_workbook(io.BytesIO(write_submission(result, s.submission_dukhan)))
    count = lambda name: sum(1 for row in sub[name].iter_rows(min_row=6) if any(c.value is not None for c in row))  # noqa: E731
    assert (count("BOM HEADER"), count("BOM PARTS"), count("MATERIAL MASTER")) == (1, 3, 4)


def test_identical_spare_under_two_versions_of_an_equipment_is_listed_once():
    from app.services.excel_writer import unique_lines
    from app.services.input_reader import InputRow

    rows = []
    for i, (tag, qty) in enumerate([("PG-1", "131"), ("PG-2", "146")]):  # same equipment, other EQPT QTY
        rows += [InputRow(10 * i + 1, {"TAG NO": tag, "EQPT QTY": qty, "MANUFACTURER PART NUMBER": "LFBSPG"}),
                 InputRow(10 * i + 2, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": "1",
                                       "MANUFACTURER PART NUMBER": "GAUGE SAVER"})]
    out = unique_lines(build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {}).lines)
    # one temp number = one BOM header: the first tag's equipment line, the identical spare once under it
    assert [(ln.serial, ln.category, ln.temp_number, ln.quantity) for ln in out] == [
        (1, "B", 40001, 131), (2, "L", 500001, 1)]


def _sap_rows(spares, tag="P-101", eq_pn="PUMP-1", first=1):
    """Equipment row + spare rows (sap, part number, qty) under one tag, input rows numbered from `first`."""
    from app.services.input_reader import InputRow

    rows = [InputRow(first, {"TAG NO": tag, "EQPT QTY": "1", "MANUFACTURER PART NUMBER": eq_pn,
                             "DESCRIPTION OF PARTS": "PUMP"})]
    for i, (sap, pn, qty) in enumerate(spares, start=1):
        rows.append(InputRow(first + i, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": qty, "SAP NUMBER": sap,
                                         "MANUFACTURER PART NUMBER": pn, "DESCRIPTION OF PARTS": f"PART {pn}"}))
    return rows


def test_same_sap_and_part_number_under_one_tag_adds_the_quantities():
    from app.services.excel_writer import unique_lines

    rows = _sap_rows([("10636722", "SEAL-A", "2"), ("10636723", "BRG-1", "1"),
                      ("10636722", " seal-a ", "3"), ("10636722", "SEAL-A", "0.5")])  # case / spaces do not count
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert len(res.lines) == 5                                                       # every input line kept
    assert [w.lines for w in res.warnings if w.code == "REPEATED_SPARE"] == [[4, 5]]
    assert not [w for w in res.warnings if w.code in ("QUANTITY_DIFFERS", "QUANTITY_NOT_ADDED")]
    assert [(ln.serial, ln.temp_number, ln.mfr_part_number, ln.quantity) for ln in unique_lines(res.lines)] == [
        (1, 40001, "PUMP-1", 1), (2, 10636722, "SEAL-A", 5.5), (3, 10636723, "BRG-1", 1)]


def test_same_sap_other_part_numbers_are_separate_lines_on_every_sheet():
    from app.services.excel_writer import write_submission

    rows = _sap_rows([("10636722", "SEAL-A", "2"), ("10636722", "SEAL-B", "4"), ("10636722", "SEAL-A", "1")])
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [w.lines for w in res.warnings if w.code == "REPEATED_SPARE"] == [[4]]  # SEAL-B is not a repeat

    s = get_settings()
    bom = openpyxl.load_workbook(io.BytesIO(write_workbook(res, DUKHAN, s.template_dukhan)))["BOM WORKING TEMPLATE"]
    hdr = {str(c.value).strip(): c.column for c in bom[1] if c.value}
    got = [(bom.cell(r, hdr["MATERIAL TEMP NUMBER"]).value, bom.cell(r, hdr["MANUFACTURERS PART NUMBER"]).value,
            bom.cell(r, hdr["QUANTITY"]).value) for r in range(6, 10)]
    assert got == [(40001, "PUMP-1", 1), (10636722, "SEAL-A", 3), (10636722, "SEAL-B", 4), (None, None, None)]

    sub = openpyxl.load_workbook(io.BytesIO(write_submission(res, s.submission_dukhan)))
    count = lambda name: sum(1 for row in sub[name].iter_rows(min_row=6) if any(c.value is not None for c in row))  # noqa: E731
    assert (count("BOM HEADER"), count("BOM PARTS"), count("MATERIAL MASTER")) == (1, 2, 3)


def test_same_sap_and_part_number_is_not_added_across_tags_of_one_header():
    from app.services.excel_writer import unique_lines

    rows = (_sap_rows([("10636722", "SEAL-A", "2"), ("10636722", "SEAL-A", "1")], tag="P-101") +
            _sap_rows([("10636722", "SEAL-A", "3")], tag="P-102", first=10))       # same pump, same total
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [t.submt for t in res.tags] == [40001, 40001]
    assert not [w for w in res.warnings if w.code == "QUANTITY_DIFFERS"]
    # one BOM: the first tag's total, not 2 + 1 + 3
    assert [(ln.category, ln.temp_number, ln.quantity) for ln in unique_lines(res.lines)] == [
        ("B", 40001, 1), ("L", 10636722, 3)]

    rows = (_sap_rows([("10636722", "SEAL-A", "2"), ("10636722", "SEAL-A", "1")], tag="P-101") +
            _sap_rows([("10636722", "SEAL-A", "2")], tag="P-102", first=10))       # same pump, other total
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [w.lines for w in res.warnings if w.code == "QUANTITY_DIFFERS"] == [[11]]
    assert [ln.quantity for ln in unique_lines(res.lines)] == [1, 3]


def test_text_quantity_of_a_repeated_spare_is_not_added():
    from app.services.excel_writer import unique_lines

    rows = _sap_rows([("10636722", "SEAL-A", "2"), ("10636722", "SEAL-A", "AS REQUIRED"), ("10636722", "SEAL-A", None)])
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [w.lines for w in res.warnings if w.code == "QUANTITY_NOT_ADDED"] == [[3]]
    assert [ln.quantity for ln in unique_lines(res.lines)] == [1, 2]


def test_verification_reason_stays_with_its_own_part_number_under_one_sap():
    from app.services.excel_writer import find_verification_sheet

    rows = _sap_rows([("10636722", "SEAL-A", "2"), ("10636722", "PUMP-1", "1")])  # 2nd = the equipment's part number
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    wb = openpyxl.load_workbook(io.BytesIO(write_workbook(res, DUKHAN, get_settings().template_dukhan)))
    ws = find_verification_sheet(wb)
    hdr = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    flagged = [(ws.cell(r, hdr["SERIAL NUMBER"]).value, ws.cell(r, hdr["MANUFACTURERS PART NUMBER"]).value)
               for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value is not None]
    assert flagged == [(3, "PUMP-1")]


def test_cabinet_is_grouped_by_its_spares_not_their_quantities():
    from app.services.excel_writer import unique_lines
    from app.services.input_reader import InputRow

    def eq(n, tag, qty="1"):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": qty, "NEW DESCRIPTION OF PARTS": "LOCAL PANEL",
                            "SUPPLIER/ OCM NAME": "ESS"})

    def sp(n, tag, sap, pos, qty="2"):
        return InputRow(n, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": qty, "SAP NUMBER": sap,
                            "POSITION NUMBER": pos})

    rows = [eq(1, "LP-1", "8"), sp(2, "LP-1", "10636722", "0010"), sp(3, "LP-1", "10636723", "0020"),
            eq(4, "LP-2", "2"), sp(5, "LP-2", "10636722", "0030"), sp(6, "LP-2", "10636723", "0040"),  # same spares
            eq(7, "LP-3"), sp(8, "LP-3", "10636722", "0010"), sp(9, "LP-3", "10636723", "0020"),
            sp(10, "LP-3", "10636722", "0090"), sp(11, "LP-3", "10636723", "0100"),  # same spares, listed twice
            eq(12, "LP-4"), sp(13, "LP-4", "10636722", "0010"),                      # other spare set
            eq(14, "LP-5"), sp(15, "LP-5", "10636722", "0010", qty="5"),             # same set, other quantity
            eq(16, "LP-6"), eq(17, "LP-7")]                                          # no spares -> by tag
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {}, cabinet_quantities=False)  # old rule
    assert [(t.tag, t.submt) for t in res.tags] == [("LP-1", 40001), ("LP-2", 40001), ("LP-3", 40001),
                                                   ("LP-4", 40002), ("LP-5", 40002), ("LP-6", 40003),
                                                   ("LP-7", 40004)]
    assert [w.lines for w in res.warnings if w.code == "REPEATED_SPARE"] == [[10, 11]]
    # LP-3's repeats add up to 4 (LP-1 has 2) and LP-5 has 5 - the templates keep LP-1's quantities
    assert [w.lines for w in res.warnings if w.code == "QUANTITY_DIFFERS"] == [[8, 9, 10, 11, 15]]
    out = unique_lines(res.lines)
    # each spare once per header, with the first tag's quantity
    assert [(ln.category, ln.temp_number, ln.bom_header_number, ln.quantity) for ln in out] == [
        ("B", 40001, 40001, 8), ("L", 10636722, 40001, 2), ("L", 10636723, 40001, 2),
        ("B", 40002, 40002, 1), ("L", 10636722, 40002, 2),
        ("B", 40003, 40003, 1), ("B", 40004, 40004, 1)]


def test_cabinet_make_and_model_with_same_spares_share_a_number():
    from app.services.input_reader import InputRow

    def eq(n, tag, make="RITTAL", model="TS8"):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "NEW DESCRIPTION OF PARTS": "MARSHALLING CABINET",
                            "EQPT MAKE": make, "EQPT MODEL": model})

    def sp(n, tag, pn, qty="1"):
        return InputRow(n, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": qty, "MANUFACTURER PART NUMBER": pn})

    rows = [eq(1, "MC-1"), sp(2, "MC-1", "FUSE"),
            eq(3, "MC-2"), sp(4, "MC-2", "fuse "),                 # same make, model, spares, quantity
            eq(5, "MC-3", model="TS9"), sp(6, "MC-3", "FUSE"),     # other model
            eq(7, "MC-4"), sp(8, "MC-4", "FUSE", qty="3")]         # same spares, other quantity -> same
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {}, cabinet_quantities=False)  # old rule
    assert [t.submt for t in res.tags] == [40001, 40001, 40002, 40001]


def test_panels_with_the_same_parts_in_other_quantities_share_a_number():
    """SPIR VEN-4460-DGTYP-4-43-0309: AB' GLM RIO PANEL at 4 locations. DIYAB has FAHAHIL's parts with about
    half the quantities -> one number; KHATIYA (2 extra parts) and JALEHA (2 parts fewer) keep their own."""
    from app.services.input_reader import InputRow

    fahahil = {"1783-CMS10DP": "2", "1786-RG6": "50", "1786-BNC": "4", "D5020S": "6", "ET-498-2TX": "1"}
    panels = [("KHATIYA", ["22-4202-LP-0103", "23-4721-LP-0001"], {**fahahil, "1783-MMS10AR": "2", "1783-MMX8SA": "2"}),
              ("FAHAHIL", ["12-4020-LP-0103", "13-4032-LP-0001"], fahahil),
              ("JALEHA", ["31-4919-LP-0104"], {"1786-RG6": "50", "1786-BNC": "4", "D5020S": "1"}),
              ("DIYAB", ["32-4001-LP-0001", "32-4002-LP-0001"],
               {"1783-CMS10DP": "1", "1786-RG6": "25", "1786-BNC": "2", "D5020S": "3", "ET-498-2TX": "1"})]
    rows, n = [], 4
    for place, tags, spares in panels:
        for tag in tags:
            rows.append(InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "EQPT MAKE": "ROCKWELL AUTOMATION",
                                     "EQPT MODEL": f"AB' GLM RIO PANEL-{place}", "NEW DESCRIPTION OF PARTS": "RIO PANEL"}))
            for pn, qty in spares.items():
                n += 1
                rows.append(InputRow(n, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": qty,
                                         "MANUFACTURER PART NUMBER": pn}))
            n += 1
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {}, cabinet_quantities=False)  # old rule
    assert [(t.tag[:2], t.submt) for t in res.tags] == [("22", 40001), ("23", 40001), ("12", 40002), ("13", 40002),
                                                       ("31", 40003), ("32", 40002), ("32", 40002)]
    # the DIYAB rows with another quantity are reported; the templates keep FAHAHIL's quantities
    assert len([w for w in res.warnings if w.code == "QUANTITY_DIFFERS"][0].lines) == 2 * 4


def test_cabinet_quantities_split_the_submt_when_enabled():
    from app.services.excel_writer import unique_lines
    from app.services.input_reader import InputRow

    def eq(n, tag, qty="1"):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": qty, "NEW DESCRIPTION OF PARTS": "LOCAL PANEL",
                            "SUPPLIER/ OCM NAME": "ESS"})

    def sp(n, tag, sap, pos, qty="2"):
        return InputRow(n, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": qty, "SAP NUMBER": sap,
                            "POSITION NUMBER": pos})

    rows = [eq(1, "LP-1", "8"), sp(2, "LP-1", "10636722", "0010"), sp(3, "LP-1", "10636723", "0020"),
            eq(4, "LP-2", "2"), sp(5, "LP-2", "10636722", "0030"), sp(6, "LP-2", "10636723", "0040"),  # same spares
            eq(7, "LP-3"), sp(8, "LP-3", "10636722", "0010"), sp(9, "LP-3", "10636723", "0020"),
            sp(10, "LP-3", "10636722", "0090", qty="7"), sp(11, "LP-3", "10636723", "0100"),  # repeats, added up
            eq(12, "LP-4"), sp(13, "LP-4", "10636722", "0010"),                      # other spare set
            eq(14, "LP-5"), sp(15, "LP-5", "10636722", "0010", qty="5"),             # same set, other quantity
            eq(16, "LP-6"), sp(17, "LP-6", "10636722", "0010", qty="2.0"),           # same quantity, other spelling
            eq(18, "LP-7"), sp(19, "LP-7", "10636722", "0010", qty=None),            # no quantity
            eq(20, "LP-8"), eq(21, "LP-9")]                                          # no spares -> by tag
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})  # default: quantities compared
    assert [(t.tag, t.submt) for t in res.tags] == [("LP-1", 40001), ("LP-2", 40001), ("LP-3", 40001),
                                                   ("LP-4", 40002), ("LP-5", 40003), ("LP-6", 40002),
                                                   ("LP-7", 40004), ("LP-8", 40005), ("LP-9", 40006)]
    # only LP-3 still differs (its repeats add up to 9 and 4; the SUBMT compares the first quantities);
    # LP-5's other quantity is now its own header
    assert [w.lines for w in res.warnings if w.code == "QUANTITY_DIFFERS"] == [[8, 9, 10, 11]]
    out = unique_lines(res.lines)
    # each configuration is its own BOM with its own quantities
    assert [(ln.category, ln.temp_number, ln.bom_header_number, ln.quantity) for ln in out] == [
        ("B", 40001, 40001, 8), ("L", 10636722, 40001, 2), ("L", 10636723, 40001, 2),
        ("B", 40002, 40002, 1), ("L", 10636722, 40002, 2),
        ("B", 40003, 40003, 1), ("L", 10636722, 40003, 5),
        ("B", 40004, 40004, 1), ("L", 10636722, 40004, None),
        ("B", 40005, 40005, 1), ("B", 40006, 40006, 1)]


def test_panels_with_other_quantities_get_their_own_number_when_enabled():
    """Same RIO panels as above with CABINET_COMPARE_QUANTITIES: DIYAB (about half FAHAHIL's quantities) is split off."""
    from app.services.input_reader import InputRow

    fahahil = {"1783-CMS10DP": "2", "1786-RG6": "50", "1786-BNC": "4", "D5020S": "6", "ET-498-2TX": "1"}
    diyab = {"1783-CMS10DP": "1", "1786-RG6": "25", "1786-BNC": "2", "D5020S": "3", "ET-498-2TX": "1"}
    rows, n = [], 4
    for place, tags, spares in [("FAHAHIL", ["12-4020-LP-0103", "13-4032-LP-0001"], fahahil),
                                ("DIYAB", ["32-4001-LP-0001", "32-4002-LP-0001"], diyab)]:
        for tag in tags:
            rows.append(InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "EQPT MAKE": "ROCKWELL AUTOMATION",
                                     "EQPT MODEL": f"AB' GLM RIO PANEL-{place}", "NEW DESCRIPTION OF PARTS": "RIO PANEL"}))
            for pn, qty in spares.items():
                n += 1
                rows.append(InputRow(n, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": qty,
                                         "MANUFACTURER PART NUMBER": pn}))
            n += 1
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [t.submt for t in res.tags] == [40001, 40001, 40002, 40002]
    assert not [w for w in res.warnings if w.code == "QUANTITY_DIFFERS"]


def test_fgs_cabinets_with_same_spares_other_quantities_get_own_submt():
    """SPIR VEN-4460-DGTYP-4-43-0004-A: FGS cabinets of one make + model, the same part list everywhere (same spare
    count) but other quantities per cabinet -> one SUBMT per quantity configuration, each BOM with its own quantities."""
    from app.services.excel_writer import unique_lines
    from app.services.input_reader import InputRow

    base = {"10209/2/1": "1", "FC-TDO-1624": "1", "61-2205.0/D": "11", "61-8006.02": "17"}
    cabinets = [("14A-5227", base),
                ("12A-4020", {**base, "10209/2/1": "2", "FC-TDO-1624": "2", "61-2205.0/D": "13", "61-8006.02": "19"}),
                ("13A-4032", {**base, "10209/2/1": "2", "FC-TDO-1624": "2", "61-2205.0/D": "13", "61-8006.02": "19"}),
                ("22A-4202", {**base, "61-2205.0/D": "10", "61-8006.02": "16"}),
                ("22A-4203", {**base, "61-2205.0/D": "10", "61-8006.02": "10"}),   # one part differs
                ("23A-4722", dict(base))]                                           # same as the first
    rows, n = [], 4
    for tag, spares in cabinets:
        rows.append(InputRow(n, {"TAG NO": f"{tag}-FGP-0801", "EQPT QTY": "1", "EQPT MAKE": "HONEYWELL/EAO",
                                 "EQPT MODEL": "RITTAL CABINET TS8 SERIES",
                                 "NEW DESCRIPTION OF PARTS": "FGS SYSTEM CABINET", "SPIR NO": "VEN-4460-DGTYP-4-43-0004-REV-A"}))
        for pn, qty in spares.items():
            n += 1
            rows.append(InputRow(n, {"TAG NO": f"{tag}-FGP-0801", "QUANTITY IDENTICAL PARTS FITTED": qty,
                                     "MANUFACTURER PART NUMBER": pn}))
        n += 1
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [(t.tag[:8], t.submt) for t in res.tags] == [("14A-5227", 40001), ("12A-4020", 40002), ("13A-4032", 40002),
                                                       ("22A-4202", 40003), ("22A-4203", 40004), ("23A-4722", 40001)]
    assert all(t.herst == "HONEYWELL/EAO" and t.spir_no == "VEN-4460-DGTYP-4-43-0004" and t.rev == "A" for t in res.tags)
    assert not [w for w in res.warnings if w.code == "QUANTITY_DIFFERS"]
    # one BOM block per SUBMT, the spare quantities of its first tag
    qty = {}
    for ln in unique_lines(res.lines):
        if ln.category == "L":
            qty.setdefault(ln.bom_header_number, []).append(ln.quantity)
    assert qty == {40001: [1, 1, 11, 17], 40002: [2, 2, 13, 19], 40003: [1, 1, 10, 16], 40004: [1, 1, 10, 10]}


def test_tag_repeated_with_other_equipment_gets_its_own_bom():
    from app.services.input_reader import InputRow

    def eq(n, tag, model):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "NEW DESCRIPTION OF PARTS": "CABINET,BREATHING APPARATUS",
                            "EQPT MAKE": "IGT", "EQPT MODEL": model, "MANUFACTURER PART NUMBER": model})

    def sp(n, tag, sap):
        return InputRow(n, {"TAG NO": tag, "QUANTITY IDENTICAL PARTS FITTED": "1", "SAP NUMBER": sap})

    rows = [eq(1, "REB-1", "JB83BA"), sp(2, "REB-1", "30047643"),
            eq(3, "REB-2", "JB83BA"), sp(4, "REB-2", "30047643"),
            eq(5, "REB-3", "JB81"), sp(6, "REB-3", "10507215"),
            eq(7, "REB-2", "JB81"), sp(8, "REB-2", "10507215"),         # same tag again, other equipment
            eq(9, "REB-1", "JB83BA"), sp(10, "REB-1", "30047643")]      # same tag again, same equipment
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [(t.tag, t.typbz, t.submt) for t in res.tags] == [
        ("REB-1", "JB83BA", 40001), ("REB-2", "JB83BA", 40001), ("REB-3", "JB81", 40002), ("REB-2", "JB81", 40002)]
    assert [(ln.tag, ln.category, ln.bom_header_number, ln.input_row) for ln in res.lines if ln.tag == "REB-2"] == [
        ("REB-2", "B", 40001, 3), ("REB-2", "L", 40001, 4), ("REB-2", "B", 40002, 7), ("REB-2", "L", 40002, 8)]
    assert [w.lines for w in res.warnings if w.code == "TAG_REPEATED_OTHER_EQUIPMENT"] == [[7]]


def test_manual_valve_is_numbered_by_mesc_number():
    from app.services.bom_builder import equipment_class, mesc_number
    from app.services.input_reader import InputRow

    assert mesc_number("GATE VALVE, 2IN, CL300 MESC 74.10.21.104.1") == "7410211041"
    assert mesc_number("BALL VALVE MESC: 7410211041") == "7410211041"
    assert mesc_number("BALL VALVE 74.10.21.104 RF") == "741021104"
    assert mesc_number("BALL VALVE 2IN CL 300 PN 12345678") is None

    def eq(n, tag, desc, make="KSB", model="X1"):
        return InputRow(n, {"TAG NO": tag, "EQPT QTY": "1", "NEW DESCRIPTION OF PARTS": desc,
                            "EQPT MAKE": make, "EQPT MODEL": model})

    rows = [eq(1, "V-1", "GATE VALVE MESC 74.10.21.104.1"),
            eq(2, "V-2", "GATE VALVE MESC 74.10.21.104.1", make="ORION", model="Z9"),  # same MESC, other make
            eq(3, "V-3", "GATE VALVE MESC 74.10.21.105.1"),                            # same make, other MESC
            eq(4, "V-4", "GLOBE VALVE"),                                               # no MESC -> make + model
            eq(5, "V-5", "CONTROL VALVE", make="FISHER", model="ED")]                  # not a manual valve
    assert [equipment_class(r) for r in rows] == ["VALVE"] * 4 + ["OTHER"]
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [t.submt for t in res.tags] == [40001, 40001, 40002, 40003, 40004]
    assert [w.lines for w in res.warnings if w.code == "VALVE_NO_MESC"] == [[4]]


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_spare_with_sap_number_3_or_8_is_category_n(plant_type):
    from app.services.excel_writer import write_submission
    from app.services.input_reader import InputRow, norm_header

    def row(n, tag, sap=None, equipment=False, pn=None):
        values = {"TAG NO": tag, "SAP NUMBER": sap, "MANUFACTURER PART NUMBER": pn or f"PN-{n}"}
        values.update({"EQPT QTY": "1"} if equipment else {"QUANTITY IDENTICAL PARTS FITTED": "1"})
        return InputRow(n, values)

    rows = [row(1, "T-1", "30161704", equipment=True),   # equipment with a 3... SAP number -> stays B
            row(2, "T-1", "30159733"), row(3, "T-1", "80012204"),   # spares 3... / 8... -> N
            row(4, "T-1", "10579834"), row(5, "T-1")]               # spare 1... / no SAP -> L
    result = build_bom(rows, plant_type, "DK01", InMemoryAllocator(), {})
    assert [ln.category for ln in result.lines] == ["B", "L", "L", "L", "L"]          # internal B / L unchanged
    assert [ln.type_category for ln in result.lines] == ["B", "N", "N", "L", "L"]
    assert (result.equipment_count, result.spare_count) == (1, 4)

    s = get_settings()
    wb = openpyxl.load_workbook(io.BytesIO(write_workbook(
        result, plant_type, s.template_dukhan if plant_type == DUKHAN else s.template_other)))
    ws = wb["BOM WORKING TEMPLATE"]
    col = {norm_header(c.value): c.column for c in ws[1] if c.value}[norm_header("MATERIAL TYPEVCATEGORY")]
    assert [ws.cell(r, col).value for r in range(6, 11)] == ["B", "N", "N", "L", "L"]
    sub = openpyxl.load_workbook(io.BytesIO(write_submission(
        result, s.submission_dukhan if plant_type == DUKHAN else s.submission_other)))
    mm = sub["MATERIAL MASTER"]
    mcol = {norm_header(c.value): c.column for c in mm[1] if c.value}[norm_header("MATERIAL TYPEVCATEGORY")]
    assert [mm.cell(r, mcol).value for r in range(6, 11)] == ["B", "N", "N", "L", "L"]


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_sap_number_is_the_temp_number_in_every_plant(plant_type):
    from app.services.input_reader import InputRow

    rows = [InputRow(1, {"TAG NO": "T-1", "EQPT QTY": "1", "SAP NUMBER": "30161704"}),
            InputRow(2, {"TAG NO": "T-1", "QUANTITY IDENTICAL PARTS FITTED": "1", "SAP NUMBER": "10579834"}),
            InputRow(3, {"TAG NO": "T-1", "QUANTITY IDENTICAL PARTS FITTED": "1", "MANUFACTURER PART NUMBER": "X"}),
            InputRow(4, {"TAG NO": "T-2", "EQPT QTY": "1"})]
    lines = build_bom(rows, plant_type, "DK01", InMemoryAllocator(), {}).lines
    assert [(ln.category, ln.temp_number) for ln in lines] == [
        ("B", 30161704), ("L", 10579834), ("L", 500001),   # SAP numbers as they are; no SAP -> counter
        ("B", 40001)]                                      # the counters were not used for SAP lines
    assert lines[0].old_material_number == "QEQU-30161704" and lines[1].bom_header_number == 30161704


def test_category_n_when_the_reused_number_is_a_sap_number():
    """A spare without SAP NUMBER that reuses a SAP number from the material master (by part number)."""
    from dataclasses import replace
    from app.services.input_reader import InputRow

    rows = [InputRow(1, {"TAG NO": "T-1", "EQPT QTY": "1"}),
            InputRow(2, {"TAG NO": "T-1", "QUANTITY IDENTICAL PARTS FITTED": "1", "MANUFACTURER PART NUMBER": "X"})]
    spare = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {}).lines[1]
    assert spare.type_category == "L"                                        # counter number 500001
    assert replace(spare, temp_number=30161704).type_category == "N"         # master's SAP number 3...
    assert replace(spare, temp_number=80012204).type_category == "N"         # 8...
    assert replace(spare, temp_number=10428450).type_category == "L"         # 1... SAP number


def _reb_rows():
    """SPIR VEN-4460-DGTYP-3-43-3001: each breathing-apparatus cabinet lists itself as its spare. Three tags have
    EQPT MODEL JB81 but part number / description JB83BA, and a JB81 spare - the part number of other tags' equipment."""
    from app.services.input_reader import InputRow

    vendor = "INTERNATIONAL GULF TRADING COMPANY /SALAM PETROLEUM"
    sap = {"JB83BA": "30047643", "JB81": "10507215"}
    rows, n = [], 4
    for tag, model, eq_pn, spare_pn, pos in [
            ("46-REB-0001", "JB83BA", "JB83BA", "JB83BA", "0010"), ("46-REB-0002", "JB83BA", "JB83BA", "JB83BA", "0010"),
            ("24-REB-0003", "JB81", "JB81", "JB81", "0010"), ("23-REB-0001", "JB81", "JB83BA", "JB81", "0020"),
            ("12-REB-0002", "JB81", "JB83BA", "JB81", "0020"), ("14-REB-0003", "JB81", "JB81", "JB81", "0010"),
            ("31-REB-0001", "JB81", "JB83BA", "JB81", "0020")]:
        rows += [InputRow(n, {"TAG NO": tag, "EQPT MAKE": vendor, "EQPT MODEL": model, "EQPT QTY": "1",
                              "MANUFACTURER PART NUMBER": eq_pn,
                              "DESCRIPTION OF PARTS": f"CABINET,BREATHING APPARATUS,SCNTN,{eq_pn}"}),
                 InputRow(n + 1, {"TAG NO": tag, "EQPT MAKE": vendor, "EQPT MODEL": model,
                                  "QUANTITY IDENTICAL PARTS FITTED": "1", "POSITION NUMBER": pos,
                                  "MANUFACTURER PART NUMBER": spare_pn, "SUPPLIER/ OCM NAME": "JOBIRD",
                                  "SAP NUMBER": sap[spare_pn],
                                  "NEW DESCRIPTION OF PARTS": "BREATHING APPARATUS CABINET"})]
        n += 2
    rows += [InputRow(n, {"TAG NO": "PG-1", "EQPT QTY": "1", "EQPT MAKE": "WIKA", "EQPT MODEL": "232.50",
                          "MANUFACTURER PART NUMBER": "232.50.100", "NEW DESCRIPTION OF PARTS": "PRESSURE GAUGE"}),
             InputRow(n + 1, {"TAG NO": "PG-1", "QUANTITY IDENTICAL PARTS FITTED": "1", "SAP NUMBER": "10000001",
                              "MANUFACTURER PART NUMBER": "gs-1", "NEW DESCRIPTION OF PARTS": "GAUGE SAVER"})]
    return rows


def test_spare_with_an_equipment_part_number_is_flagged():
    res = build_bom(_reb_rows(), DUKHAN, "DK01", InMemoryAllocator(), {})
    assert not [ln for ln in res.lines if ln.category == "B" and ln.review]       # equipment lines never
    spares = {ln.tag: ln.review for ln in res.lines if ln.category == "L"}
    for tag in ("46-REB-0001", "46-REB-0002", "24-REB-0003", "14-REB-0003"):      # its own equipment's PN
        assert spares[tag] == [f"Same part number as its equipment (tag {tag}) - check whether this is the "
                               "equipment itself, not a spare."]
    for tag in ("23-REB-0001", "12-REB-0002", "31-REB-0001"):                     # only other equipment's PN
        assert spares[tag] == [f"Part number JB81 is the equipment part number of tag(s) 24-REB-0003, 14-REB-0003, "
                               f"but its own equipment (tag {tag}) has JB83BA - check which is right."]
    assert spares["PG-1"] == []                                                    # no equipment has GS-1
    assert [w.lines for w in res.warnings if w.code == "VERIFY"] == [[2, 4, 6, 8, 10, 12, 14]]
    # the BOM itself is unchanged by the review
    assert [(ln.tag, ln.category, ln.temp_number) for ln in res.lines][:4] == [
        ("46-REB-0001", "B", 40001), ("46-REB-0001", "L", 30047643),
        ("46-REB-0002", "B", 40001), ("46-REB-0002", "L", 30047643)]


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_verification_sheet_lists_only_flagged_lines(plant_type):
    from app.services.combiner import WORKING, combine

    s = get_settings()
    res = build_bom(_reb_rows(), plant_type, "DK01", InMemoryAllocator(), {})
    content = write_workbook(res, plant_type, s.template_dukhan if plant_type == DUKHAN else s.template_other)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert [ws.title.strip() for ws in wb.worksheets][2:] == ["VERIFICATION"]
    assert wb.active.title.strip() == "BOM WORKING TEMPLATE"           # still opens on the BOM sheet
    # BOM sheet: 40001 + its spare, 40002 (all JB81 cabinets, first tag 24-REB-0003) + its spare, the gauge + its spare
    assert sum(1 for r in wb.worksheets[0].iter_rows(min_row=6, values_only=True) if any(r)) == 6
    ws = wb["VERIFICATION"]
    hdr = {c.value: c.column for c in ws[1]}
    data = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
    assert [(r[hdr["SERIAL NUMBER"] - 1], r[hdr["TAG NO"] - 1], r[hdr["MATERIAL TEMP NUMBER"] - 1]) for r in data] == [
        (2, "46-REB-0001", 30047643), (4, "24-REB-0003", 10507215)]
    # the JB81 spare is listed once (under 24-REB-0003); the conflicting tags' reasons come with it
    reason = data[1][hdr["CHECK REASON"] - 1]
    for tag, row in (("23-REB-0001", 11), ("12-REB-0002", 13), ("31-REB-0001", 17)):
        assert f"(tag {tag}, input row {row})" in reason and "check which is right" in reason
    assert ws.cell(2, hdr["SERIAL NUMBER"]).hyperlink.location.endswith("!A7")  # row of serial 2 on the BOM sheet
    assert ws.freeze_panes == "A2"

    # combined: the second file's serials (and links) follow the first file's BOM rows
    vs = openpyxl.load_workbook(io.BytesIO(combine([content, content], WORKING)))["VERIFICATION"]
    assert [r[0].value for r in vs.iter_rows(min_row=2) if r[0].value is not None] == [2, 4, 8, 10]
    assert vs.cell(4, 1).hyperlink.location.endswith("!A13")


@pytest.mark.parametrize("plant_type", [DUKHAN, OTHER])
def test_no_verification_sheet_without_flagged_items(sample_bytes, plant_type):
    from dataclasses import replace

    from app.services.bom_builder import BomResult
    from app.services.combiner import WORKING, combine

    s = get_settings()
    template = s.template_dukhan if plant_type == DUKHAN else s.template_other

    def sheets(content):
        return [ws.title.strip() for ws in openpyxl.load_workbook(io.BytesIO(content)).worksheets]

    plain = write_workbook(build(sample_bytes, plant_type), plant_type, template)    # nothing flagged
    assert sheets(plain) == ["BOM WORKING TEMPLATE", "SPIR TAG vs SUBMT"]
    assert not [w for w in build(sample_bytes, plant_type).warnings if w.code == "VERIFY"]

    # flagged lines without tag details do not count either
    res = build_bom(_reb_rows(), plant_type, "DK01", InMemoryAllocator(), {})
    no_tags = BomResult([replace(ln, tag=" ") for ln in res.lines], res.tags, res.warnings, res.spir_numbers)
    assert sheets(write_workbook(no_tags, plant_type, template)) == ["BOM WORKING TEMPLATE", "SPIR TAG vs SUBMT"]

    # combined: the sheet only when a file has flagged items; their serials follow the BOM rows before them
    flagged = write_workbook(res, plant_type, template)
    assert sheets(combine([plain, plain], WORKING)) == ["BOM WORKING TEMPLATE", "SPIR TAG vs SUBMT"]
    combined = openpyxl.load_workbook(io.BytesIO(combine([plain, flagged], WORKING)))
    assert combined.active.title.strip() == "BOM WORKING TEMPLATE"
    plain_rows = sum(1 for r in openpyxl.load_workbook(io.BytesIO(plain)).worksheets[0]
                     .iter_rows(min_row=6, values_only=True) if any(r))
    vs = combined["VERIFICATION"]
    assert [r[0].value for r in vs.iter_rows(min_row=2) if r[0].value is not None] == [plain_rows + 2, plain_rows + 4]


def test_base_model_drops_a_location_suffix_only():
    from app.services.bom_builder import base_model

    for model in ("AB' GLM RIO PANEL-KHATIYA", "AB' GLM RIO PANEL-FAHAHIL", "ab' glm rio panel - jaleha",
                  "AB' GLM RIO PANEL-DIYAB", "AB' GLM RIO PANEL"):
        assert base_model(model) == "AB' GLM RIO PANEL"
    # model codes keep their dashes: digits in the suffix, no space, base ending in a number, short suffix
    for model, kept in (("NFF-SPH-4W", "NFF-SPH-4W"), ("NI-GVO-300FF", "NI-GVO-300FF"), ("BRK 703-704", "BRK 703-704"),
                        ("DRAGER PAS LITE-P/N: 3357774", "DRAGER PAS LITE-P/N: 3357774"), ("TS8-ABCD", "TS8-ABCD"),
                        ("ROSEMOUNT 3051-HART", "ROSEMOUNT 3051-HART"), ("GLM PANEL-FF", "GLM PANEL-FF"),
                        ("JB83BA", "JB83BA"), ("232.50", "232.50")):
        assert base_model(model) == kept


@pytest.mark.parametrize("description", ["GAS DETECTOR", "LOCAL CONTROL PANEL"])  # other equipment / cabinet
def test_model_with_a_location_suffix_shares_the_equipment_number(description):
    from app.services.input_reader import InputRow

    places = ["KHATIYA", "FAHAHIL", "JALEHA", "DIYAB"]
    rows = []
    for i, place in enumerate(places):
        rows += [InputRow(10 * i + 1, {"TAG NO": f"GLM-{i}", "EQPT QTY": "1", "EQPT MAKE": "ABB",
                                       "EQPT MODEL": f"AB' GLM RIO PANEL-{place}", "NEW DESCRIPTION OF PARTS": description}),
                 InputRow(10 * i + 2, {"TAG NO": f"GLM-{i}", "QUANTITY IDENTICAL PARTS FITTED": "1",
                                       "SAP NUMBER": "10636722"})]
    rows += [InputRow(50, {"TAG NO": "GLM-X", "EQPT QTY": "1", "EQPT MAKE": "ABB", "EQPT MODEL": "AB' GLM RIO PANEL-2",
                           "NEW DESCRIPTION OF PARTS": description}),  # digit suffix -> another model
             InputRow(51, {"TAG NO": "GLM-X", "QUANTITY IDENTICAL PARTS FITTED": "1", "SAP NUMBER": "10636722"})]
    res = build_bom(rows, DUKHAN, "DK01", InMemoryAllocator(), {})
    assert [t.submt for t in res.tags] == [40001] * 4 + [40002]
    # the model shown in the output is never changed
    assert [t.typbz for t in res.tags] == [f"AB' GLM RIO PANEL-{p}" for p in places] + ["AB' GLM RIO PANEL-2"]
