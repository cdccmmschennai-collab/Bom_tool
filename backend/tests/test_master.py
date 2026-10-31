"""Material master (material_master_import): continue numbering, reuse numbers, save new lines.

Runs against PostgreSQL (TEST_DATABASE_URL, wiped by these tests).
"""
import pytest
from sqlalchemy import func, insert, select
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from app.database import Base, engine

try:
    with engine.connect():
        pass
except OperationalError:  # pragma: no cover
    pytest.skip("PostgreSQL test database not reachable", allow_module_level=True)

from app.models import material_master_import as master  # noqa: E402
from app.services.bom_builder import DUKHAN, InMemoryAllocator, build_bom  # noqa: E402
from app.services.input_reader import read_input  # noqa: E402
from app.services.material_master import save_to_master  # noqa: E402
from app.services.numbering import DbAllocator  # noqa: E402


@pytest.fixture
def db():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        yield s


def add_master(db, plant, temp, equipment, sap=None, spf=None, pn=None, project="P1"):
    db.execute(insert(master).values(
        material_temp_number=temp, maintenance_planning_plant=plant, material_is_equipment_indicator=equipment,
        material_type_category="B" if equipment else "L", sap_material_number=sap,
        old_material_number_spf_number=spf, manufacturers_part_number=pn, project_number=project))


def test_numbering_continues_after_master(db):
    add_master(db, "T1", 40010, True)
    add_master(db, "T1", 500020, False)
    add_master(db, "T1", 80012771, False, sap="80012771")  # SAP number in the temp column - ignored
    add_master(db, "T2", 40500, True)                       # other plant - counts too (no duplicates)
    add_master(db, "T2", 40500, True)                       # duplicate value - ignored
    add_master(db, "T2", 500030, False)
    alloc = DbAllocator(db, "T1", "P1")
    assert alloc.equipment("TAG:NEW") == 40501
    assert alloc.spare("DESC:NEW") == 500031


def test_spare_reuses_master_number(db):
    add_master(db, "T1", 500005, False, sap="123")
    add_master(db, "T1", 500006, False, spf="ABC-1")
    add_master(db, "T2", 500007, False, spf="XYZ")
    add_master(db, "T1", 500008, False, spf="PRJ2", project="P2")
    add_master(db, "T1", 500009, False, spf="NOPRJ", project=None)  # imported before projects existed
    alloc = DbAllocator(db, "T1", " p1 ")                           # case and spaces of the project ignored
    assert alloc.spare("SAP:123") == 500005
    assert alloc.spare("SPF:ABC-1") == 500006
    assert alloc.spare("SPF:XYZ") == 500010    # only found for plant T2 -> new number after the highest
    assert alloc.spare("SPF:PRJ2") == 500011   # same plant, other project -> new
    assert alloc.spare("SPF:NOPRJ") == 500012  # same plant, no project -> new
    assert alloc.new_numbers == 3


def test_save_to_master_adds_new_lines_once(db, sample_bytes):
    result = build_bom(read_input(sample_bytes), DUKHAN, "T1", InMemoryAllocator(), {})
    temps = {line.temp_number for line in result.lines}
    assert save_to_master(db, result, "T1", "P1") == len(temps)
    assert save_to_master(db, result, "T1", "p1") == 0  # same plant + project again -> nothing new
    assert save_to_master(db, result, "T1", "P2") == len(temps)  # another project has its own rows
    count = db.execute(select(func.count()).select_from(master)
                       .where(master.c.maintenance_planning_plant == "T1", master.c.project_number == "P1")).scalar()
    assert count == len(temps)
    eq = db.execute(select(master).where(master.c.material_temp_number == 40001, master.c.project_number == "P1")).one()
    assert eq.material_type_category == "B" and eq.material_is_equipment_indicator is True


def test_spare_reuses_master_number_by_part_number(db):
    db.execute(insert(master).values(material_temp_number=500010, maintenance_planning_plant="T1", project_number="P1",
                                     material_is_equipment_indicator=False, manufacturers_part_number="1SBL  137001r1110"))
    add_master(db, "T1", 500020, False)
    alloc = DbAllocator(db, "T1", "P1")
    assert alloc.spare("PN:1SBL 137001R1110") == 500010
    assert alloc.spare("PN:NEW-PART") == 500021
    db.commit()
    again = DbAllocator(db, "T1", "P1")  # next upload
    assert again.spare("PN:NEW-PART") == 500021
    assert again.spare("PN:OTHER-PART") == 500022


def test_grille_reuses_master_number_by_part_number_and_description(db):
    supply = "SUPPLY GRILLE IN ALUMINIUM WITH DOUBLE BLADES WITH ADJUSTABLE DAMPER-600X400MM"
    ret = "RETURN GRILLE IN ALUMINIUM WITH DOUBLE BLADES WITH ADJUSTABLE DAMPER-600X500MM"
    db.execute(insert(master).values(material_temp_number=500030, maintenance_planning_plant="T1", project_number="P1",
                                     material_is_equipment_indicator=False, manufacturers_part_number="GR-1",
                                     material_description="Supply grille in aluminium with double blades with "
                                                          "adjustable damper – 600 x 400 mm"))
    alloc = DbAllocator(db, "T1", "P1")
    assert alloc.spare(f"PNDESC:GR-1|{supply}") == 500030      # same part number + description (case, dash ignored)
    assert alloc.spare(f"PNDESC:GR-1|{ret}") == 500031         # same part number, other description -> new
    assert alloc.new_numbers == 1

def test_numbers_are_reused_only_for_the_same_plant_and_project(db):
    keys = [("equipment", "PN:JB82"), ("equipment", "MESC:7410211041"), ("equipment", "MM:ABB|RIO PANEL"),
            ("spare", "PN:O-RING 1"), ("spare", "SPF:ABC-1")]

    def numbers(plant, project):
        alloc = DbAllocator(db, plant, project)
        return [getattr(alloc, kind)(key) for kind, key in keys]

    first = numbers("T1", "P1")
    db.commit()
    assert numbers("T1", "p1") == first                          # same plant + project -> reused
    other_project = numbers("T1", "P2")                          # same plant, other project -> new
    other_plant = numbers("T2", "P1")                            # other plant, same project -> new
    every = first + other_project + other_plant
    assert len(set(every)) == len(every)                         # a number is never issued twice


def test_part_number_in_another_plants_master_is_not_reused(db):
    add_master(db, "T2", 40100, True, pn=" ex-cell  t3 ")       # equipment row of plant T2
    add_master(db, "T2", 500050, False, pn="P-1")               # spare row of plant T2
    add_master(db, "T1", 500060, False, pn="P-2", project="P9")  # same plant, other project
    alloc = DbAllocator(db, "T1", "P1")
    assert alloc.equipment("PN:EX-CELL T3") == 40101
    assert alloc.spare("PN:P-1") == 500061
    assert alloc.spare("PN:P-2") == 500062
    assert alloc.new_numbers == 3


def test_sap_number_of_another_plant_is_not_reused(db):
    add_master(db, "T2", 80012204, False, pn="SAP-PN")          # Other plant: SAP number in the temp column
    add_master(db, "T1", 80012999, False, pn="SAME-PLANT")      # same plant: reused as before
    alloc = DbAllocator(db, "T1", "P1")
    assert alloc.spare("PN:SAP-PN") == 500001
    assert alloc.spare("PN:SAME-PLANT") == 80012999


def test_equipment_and_spare_numbers_stay_apart(db):
    add_master(db, "T1", 500060, False, pn="BOTH")              # the part is known as a spare only
    alloc = DbAllocator(db, "T1", "P1")
    assert alloc.equipment("PN:BOTH") == 40001


def test_tag_numbers_stay_per_plant_and_project(db):
    t1 = DbAllocator(db, "T1", "P1").equipment("TAG:34-QD-23")
    assert DbAllocator(db, "T1", "P1").equipment("TAG:34-QD-23") == t1
    assert DbAllocator(db, "T2", "P1").equipment("TAG:34-QD-23") != t1
    assert DbAllocator(db, "T1", "P2").equipment("TAG:34-QD-23") != t1


def test_spare_never_takes_an_equipment_number(db):
    add_master(db, "T1", 40100, True, pn="JB82")                # the part is known as an equipment only
    alloc = DbAllocator(db, "T1", "P1")
    assert alloc.spare("PN:JB82") == 500001                     # a spare gets a spare number of its own
    assert alloc.equipment("PN:JB82") == 40100                  # the equipment keeps its number
    assert alloc.spare("PN:JB82") == 500001 and alloc.new_numbers == 1   # repeated: reused, not new


def test_templates_take_the_master_description_of_matching_spares(db):
    from app.services.input_reader import InputRow
    from app.services.material_master import with_master_descriptions

    def spare(n, sap, spf, desc):
        return InputRow(n, {"TAG NO": "T-1", "QUANTITY IDENTICAL PARTS FITTED": "1", "SAP NUMBER": sap,
                            "OLD MATERIAL NUMBER/SPF NUMBER": spf, "DESCRIPTION OF PARTS": desc})

    rows = [InputRow(1, {"TAG NO": "T-1", "EQPT QTY": "1", "DESCRIPTION OF PARTS": "SPIR EQUIPMENT"}),
            spare(2, "30159733", "4460-2433001-01L01", "SPIR A"),   # number + SPF match (same plant preferred)
            spare(3, "30157603", "4460-2433001-01l02 ", "SPIR B"),  # match ignoring case / spaces, other plant only
            spare(4, "30157604", "4460-2433001-01L03", "SPIR C"),   # number matches, SPF does not
            spare(5, "30157605", "4460-2433001-01L04", "SPIR D"),   # match, but the master has no description
            spare(6, "30157606", None, "SPIR E")]                   # no SPF number
    result = build_bom(rows, DUKHAN, "T1", InMemoryAllocator(), {})
    db.execute(insert(master), [
        dict(material_temp_number=40001, maintenance_planning_plant="T1", material_description="MASTER EQUIPMENT",
             old_material_number_spf_number="QEQU-40001"),                       # equipment: never used
        dict(material_temp_number=30159733, maintenance_planning_plant="T2", material_description="OTHER PLANT",
             old_material_number_spf_number="4460-2433001-01L01"),
        dict(material_temp_number=30159733, maintenance_planning_plant="T1", material_description="SAME PLANT",
             old_material_number_spf_number="4460-2433001-01L01"),
        dict(material_temp_number=30157603, maintenance_planning_plant="T2", material_description="ONLY OTHER PLANT",
             old_material_number_spf_number="4460-2433001-01L02"),
        dict(material_temp_number=30157604, maintenance_planning_plant="T1", material_description="WRONG SPF",
             old_material_number_spf_number="4460-2-43-3001-L03"),               # other format: not the same text
        dict(material_temp_number=30157605, maintenance_planning_plant="T1", material_description=" ",
             old_material_number_spf_number="4460-2433001-01L04"),
        dict(material_temp_number=30157606, maintenance_planning_plant="T1", material_description="NO SPF IN SPIR",
             old_material_number_spf_number="X"),
    ])
    templates = with_master_descriptions(db, result, "T1")
    assert [ln.description for ln in templates.lines] == [
        "SPIR EQUIPMENT", "SAME PLANT", "ONLY OTHER PLANT", "SPIR C", "SPIR D", "SPIR E"]
    assert [ln.description for ln in result.lines] == [                       # the result itself is unchanged
        "SPIR EQUIPMENT", "SPIR A", "SPIR B", "SPIR C", "SPIR D", "SPIR E"]


def test_conversion_writes_master_description_to_both_templates(db, sample_bytes):
    import io

    import openpyxl

    from app.models import ExtractedLine, Plant
    from app.services.converter import convert
    from app.services.input_reader import norm_header

    spf = next(ln.old_material_number for ln in build_bom(read_input(sample_bytes), DUKHAN, None, InMemoryAllocator(),
                                                          {}).lines if ln.temp_number == 10231307)
    db.execute(insert(master).values(material_temp_number=10231307, maintenance_planning_plant=None,
                                     material_description="MASTER DESCRIPTION", old_material_number_spf_number=spf))
    plant = Plant(plant_type="DUKHAN", code="T1", name="Test plant")
    db.add(plant)
    db.commit()
    job = convert(db, plant_type="DUKHAN", plant=plant, content=sample_bytes, filename="in.xlsx", user_name="t",
                  project_number="P1")

    def descriptions(ws, first=6):
        col = {norm_header(c.value): c.column for c in ws[1] if c.value}
        rows = [r for r in range(first, ws.max_row + 1) if ws.cell(r, col[norm_header("MATERIAL TEMP NUMBER")]).value]
        return {ws.cell(r, col[norm_header("MATERIAL TEMP NUMBER")]).value:
                ws.cell(r, col[norm_header("MATERIAL DESCRIPTION")]).value for r in rows}

    working = openpyxl.load_workbook(io.BytesIO(job.output_file))["BOM WORKING TEMPLATE"]
    submission = openpyxl.load_workbook(io.BytesIO(job.submission_file))
    assert descriptions(working)[10231307] == "MASTER DESCRIPTION"
    assert descriptions(submission["BOM PARTS"])[10231307] == "MASTER DESCRIPTION"
    assert descriptions(submission["MATERIAL MASTER"])[10231307] == "MASTER DESCRIPTION"
    assert descriptions(working)[40001] == "LV SWITCHGEAR"  # equipment: the SPIR description
    spir = db.scalar(select(ExtractedLine.description).where(ExtractedLine.temp_number == "10231307").limit(1))
    assert spir != "MASTER DESCRIPTION"                                       # Parts Master keeps the SPIR text
