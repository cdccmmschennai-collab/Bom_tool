"""Write a BomResult into the Dukhan / Other-plant template.

The template file is loaded as-is and cell VALUES are written, so header colours, fonts,
borders and widths stay exactly as in the template. No fills or conditional formatting are added.
Every cell of the filled data table (all header columns x all data rows, empty or not) gets a thin
border (like the template's header cells), and the header block is frozen so it stays visible while
scrolling. All text in the data table is black (the template pre-colours a few LEN() cells red).
"""
from __future__ import annotations

import io
from copy import copy
from dataclasses import replace
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from .bom_builder import DUKHAN, BomLine, BomResult
from .input_reader import norm_header

BOM_SHEET = "BOM WORKING TEMPLATE"
TAG_SHEET = "SPIR TAG vs SUBMT"
BOM_FIRST_DATA_ROW = 6   # rows 1-5 are the template header block
TAG_FIRST_DATA_ROW = 3   # rows 1-2 are the template header block

_THIN = Side(style="thin")
VALUE_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
BLACK = "FF000000"

# template header text -> BomLine attribute ("=LEN:<header>" writes a live LEN() formula)
BOM_COLUMNS: dict[str, str] = {
    "SERIAL NUMBER": "serial",
    "SAP MATERIAL NUMBER": "sap_material_number",          # Dukhan only
    "MATERIAL TEMP NUMBER": "temp_number",
    "MATERIAL TYPEVCATEGORY": "type_category",  # N for a spare with a 3/8 SAP number
    "OLD MATERIAL NUMBER/SPF NUMBER": "old_material_number",
    "LEN(OLD MAT/SPF NUMBER = 18)": "=LEN:OLD MATERIAL NUMBER/SPF NUMBER",
    "MATERIAL DESCRIPTION": "description",
    "LEN(MAT.DESC = 40)": "=LEN:MATERIAL DESCRIPTION",
    "MAINTENANCE PLANNING PLANT": "planning_plant",
    "BASE UNIT OF MEASURE(T006)": "uom",
    "MANUFACTURERS PART NUMBER": "mfr_part_number",
    "LEN(MFR PART NUMBER = 18)": "=LEN:MANUFACTURERS PART NUMBER",
    "MANUFACTURER NAME": "mfr_name",
    "LEN(MANUFACTURER=30)": "=LEN:MANUFACTURER NAME",
    "MATERIAL IS AN EQUIPMENT INDICATOR": "equipment_indicator",
    "BOM HEADER NUMBER": "bom_header_number",
    "QUANTITY": "quantity",
    "POSITION NUMBER": "position_number",
    "CDC_SPIR NUMBER": "spir_number_full",  # with revision (also on the submission); the tag sheet and the master keep the base
    "DELIVERY TIME IN WEEKS": "delivery_weeks",
    "MANUFACTURER COUNTRY CODE": "country_code",
    "MM REQUIREMENT MOVING AVERAGE PRICE": "price",
    "CDC_ADD INFORMATION": "add_information_text",
    "REORDER POINT (MINIMUM QTY)": "min_qty",
    "MAXIMUM STOCK LEVEL(MAX QTY)": "max_qty",
    "MANUFACTURER COUNTRY NAME": "country_name",
    "CURRENCY CODE REF": "currency",
}
# Columns that the logic says to leave empty are simply not written:
# EXTERNAL NUMBER ASSIGNED FOR ASSEMBLY, MESC/SPF_REF, SAFETY STOCK, CRITICALITY, MM MATERIAL TYPE,
# MM MATERIAL GROUP, PO TEXT, CDC_SAP CODE REMARKS, CDC_CONSTRUCTION TYPE STATUS, ASSEMBLY,
# BOM HEADER-REMARKS, PRD / QC columns.

TAG_COLUMNS = ["EQFNR", "SUBMT", "SPIR NO", "REV", "HERST", "TYPBZ", "MAPAR", "SERGE", "SPIR TYPE", "REMARKS"]
TAG_ATTRS = ["tag", "submt", "spir_no", "rev", "herst", "typbz", "mapar", "serge", "spir_type", None]

# Third sheet of the working template (not part of the template file): only the lines the tool is not sure
# about, with the reason(s), for a manual check. header -> BomLine attribute; "=REASON" = the reasons,
# None = left empty for the checker. SERIAL NUMBER links to the line on the BOM sheet.
VERIFY_SHEET = "VERIFICATION"
VERIFY_FIRST_DATA_ROW = 2  # row 1 = headers
VERIFY_COLUMNS: dict[str, str | None] = {
    "SERIAL NUMBER": "serial",
    "TAG NO": "tag",
    "INPUT EXCEL ROW": "input_row",
    "MATERIAL TYPEVCATEGORY": "type_category",
    "MATERIAL TEMP NUMBER": "temp_number",
    "BOM HEADER NUMBER": "bom_header_number",
    "OLD MATERIAL NUMBER/SPF NUMBER": "old_material_number",
    "MATERIAL DESCRIPTION": "description",
    "MANUFACTURERS PART NUMBER": "mfr_part_number",
    "MANUFACTURER NAME": "mfr_name",
    "QUANTITY": "quantity",
    "CHECK REASON": "=REASON",
    "CHECKED BY": None,
    "CHECK RESULT / ACTION": None,
}
VERIFY_WIDTHS = {"MATERIAL DESCRIPTION": 45, "CHECK REASON": 70, "CHECK RESULT / ACTION": 30,
                 "OLD MATERIAL NUMBER/SPF NUMBER": 22, "MANUFACTURERS PART NUMBER": 22, "MANUFACTURER NAME": 26}


class TemplateError(RuntimeError):
    pass


# ---------------------------------------------------------------- one row per unique record
def _row_key(line: BomLine, attrs: list[str]) -> tuple:
    return tuple(getattr(line, a) for a in attrs)


def _value_attrs(columns: dict[str, str]) -> list[str]:
    """Attributes a sheet writes, without SERIAL NUMBER and the LEN() formulas (they follow the values)."""
    return [a for a in columns.values() if a != "serial" and not a.startswith("=LEN:")]


def unique_lines(lines: list[BomLine]) -> list[BomLine]:
    """Lines of the working template, each unique record once. The sheet has no tag column, so tags that
    share an equipment temp number are one BOM: its block (B line of the first such tag + its spares) is
    listed once, and a spare a later tag adds that is not listed yet goes under that block.
    Each spare (temp number) is listed once per BOM header - its first line, so the first tag's quantity
    and position. SERIAL NUMBER is renumbered 1..n; the lines themselves are not changed."""
    blocks: dict[object, list[BomLine]] = {}
    seen: set[tuple] = set()  # (BOM header, spare temp number) already listed
    block: object = None
    for line in lines:
        key = (line.bom_header_number, line.temp_number)
        if line.category == "B" or not blocks:
            block = line.temp_number
            if block not in blocks:
                blocks[block] = [line]
            continue
        if key not in seen:
            seen.add(key)
            blocks[block].append(line)
    return [replace(line, serial=i) for i, line in enumerate((ln for b in blocks.values() for ln in b), start=1)]


def _unique_rows(lines: list[BomLine], attrs: list[str]) -> list[BomLine]:
    """Drop lines whose values in these columns repeat an earlier line (first one kept, order kept)."""
    seen: set[tuple] = set()
    out = []
    for line in lines:
        key = _row_key(line, attrs)
        if key not in seen:
            seen.add(key)
            out.append(line)
    return out


def _put(ws: Worksheet, row: int, col: int, value) -> None:
    ws.cell(row, col, value)


def _border_table(ws: Worksheet, first_row: int, row_count: int) -> None:
    """Border every cell of the data table - all header columns of every data row, empty cells
    included. The template pre-formats a few blank rows below with borders; those are dropped so
    the grid ends at the last data row. Only the border is touched."""
    last_row = first_row + row_count - 1
    last_col = max(_header_index(ws).values())
    for row in ws.iter_rows(min_row=first_row, max_row=max(ws.max_row, last_row)):
        for c in row:
            if c.row <= last_row and c.column <= last_col:
                c.border = VALUE_BORDER
            elif c.has_style:
                c.border = Border()


def _black_text(ws: Worksheet, first_row: int, row_count: int) -> None:
    """Make the text of every data-table cell black. Only the font colour changes - font name, size,
    bold etc. stay as in the template. The header block is not touched."""
    last_col = max(_header_index(ws).values())
    for row in ws.iter_rows(min_row=first_row, max_row=first_row + row_count - 1, max_col=last_col):
        for c in row:
            color = c.font.color
            if color is None or (color.type == "rgb" and color.rgb == BLACK):
                continue  # already black (no colour = Excel's automatic black)
            font = copy(c.font)
            font.color = Color(rgb=BLACK)
            c.font = font


def _sheet(wb: openpyxl.Workbook, name: str) -> Worksheet:
    for ws in wb.worksheets:
        if ws.title.strip().upper() == name.upper():
            return ws
    raise TemplateError(f"Sheet '{name}' not found in the template.")


def set_active_sheet(wb: openpyxl.Workbook, ws: Worksheet) -> None:
    """Open the workbook on `ws` with only that tab selected.

    `wb.active = ...` does not unselect the tab the template was saved on, and two selected tabs make
    Excel open the file in [Group] mode - which disables Filter and Sort.
    """
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = sheet is ws
    wb.active = wb.worksheets.index(ws)


def _header_index(ws: Worksheet, row: int = 1) -> dict[str, int]:
    return {norm_header(c.value): c.column for c in ws[row] if c.value is not None}


def _ensure_spir_type_column(ws: Worksheet) -> None:
    """The template's tag sheet has no SPIR TYPE column - insert it before REMARKS,
    copying the header style of the neighbouring column (no new colours)."""
    idx = _header_index(ws)
    if norm_header("SPIR TYPE") in idx:
        return
    remarks = idx.get(norm_header("REMARKS"))
    if remarks is None:
        remarks = ws.max_column + 1
    else:
        ws.insert_cols(remarks)
        # shift column widths right (openpyxl does not move them)
        for col in range(ws.max_column, remarks, -1):
            src = ws.column_dimensions[get_column_letter(col - 1)]
            ws.column_dimensions[get_column_letter(col)].width = src.width
    for r, text in ((1, "SPIR TYPE"), (2, "SPIR TYPE")):
        cell = ws.cell(r, remarks, text)
        cell._style = copy(ws.cell(r, remarks - 1)._style)
    ws.column_dimensions[get_column_letter(remarks)].width = max(
        ws.column_dimensions[get_column_letter(remarks - 1)].width or 10, 26)


def find_verification_sheet(wb: openpyxl.Workbook) -> Worksheet | None:
    """The VERIFICATION sheet of a working template, None when the file has none (nothing flagged, or made
    before the sheet existed)."""
    return next((ws for ws in wb.worksheets if ws.title.strip().upper() == VERIFY_SHEET), None)


def verification_sheet(wb: openpyxl.Workbook) -> Worksheet:
    """The VERIFICATION sheet of a working template; added with its header row when the file has none."""
    ws = find_verification_sheet(wb)
    if ws is not None:
        return ws
    ws = wb.create_sheet(VERIFY_SHEET)
    for col, header in enumerate(VERIFY_COLUMNS, start=1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True)
        cell.border = VALUE_BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = VERIFY_WIDTHS.get(header, 16)
    return ws


def link_serial(cell, bom: Worksheet, serial: int) -> None:
    """SERIAL NUMBER cell of the VERIFICATION sheet: the serial, linked to its line on the BOM sheet."""
    cell.value = serial
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{bom.title}'!A{BOM_FIRST_DATA_ROW + serial - 1}")


def finish_verification(ws: Worksheet, row_count: int) -> None:
    """Borders, wrapped reasons, frozen header row and a filter - like the other sheets, no colours."""
    _border_table(ws, VERIFY_FIRST_DATA_ROW, row_count)
    reason_col = _header_index(ws).get(norm_header("CHECK REASON"))
    if reason_col:
        for r in range(VERIFY_FIRST_DATA_ROW, VERIFY_FIRST_DATA_ROW + row_count):
            ws.cell(r, reason_col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = f"A{VERIFY_FIRST_DATA_ROW}"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{VERIFY_FIRST_DATA_ROW + row_count - 1}"


def _write_verification(wb: openpyxl.Workbook, bom: Worksheet, all_lines: list[BomLine],
                        lines: list[BomLine]) -> None:
    """Fill the VERIFICATION sheet with the lines of the BOM sheet (`lines`, serial = its row there) that have a
    reason to be checked and a TAG NO. A record listed once on the BOM sheet also collects the reasons of its
    repeats (other tags / positions, dropped by unique_lines), marked with their tag and input row.
    Nothing flagged -> the sheet is not created."""
    def key(line: BomLine) -> tuple:
        return line.category, line.bom_header_number, line.temp_number

    def has_tag(line: BomLine) -> bool:
        return bool(line.tag and str(line.tag).strip())

    kept = {key(line): line for line in lines}
    reasons: dict[tuple, list[str]] = {}
    for line in all_lines:
        k = key(line)
        if not line.review or not has_tag(line) or k not in kept:
            continue
        own = line.input_row == kept[k].input_row
        for text in line.review:
            text = text if own else f"{text} (tag {line.tag}, input row {line.input_row})"
            if text not in reasons.setdefault(k, []):
                reasons[k].append(text)

    flagged = [line for line in lines if key(line) in reasons and has_tag(line)]
    if not flagged:
        return
    ws = verification_sheet(wb)
    for i, line in enumerate(flagged):
        r = VERIFY_FIRST_DATA_ROW + i
        for col, attr in enumerate(VERIFY_COLUMNS.values(), start=1):
            if attr == "serial":
                link_serial(ws.cell(r, col), bom, line.serial)
            elif attr == "=REASON":
                _put(ws, r, col, "\n".join(reasons[key(line)]))
            elif attr is not None:
                value = getattr(line, attr)
                if value is not None:
                    _put(ws, r, col, value)
    finish_verification(ws, len(flagged))


def write_workbook(result: BomResult, plant_type: str, template_path: Path) -> bytes:
    wb = openpyxl.load_workbook(template_path)

    # ---- sheet 1: BOM WORKING TEMPLATE
    ws = _sheet(wb, BOM_SHEET)
    idx = _header_index(ws)
    targets: list[tuple[int, str]] = []
    for header, attr in BOM_COLUMNS.items():
        col = idx.get(norm_header(header))
        if col is None:
            if header == "SAP MATERIAL NUMBER" and plant_type != DUKHAN:
                continue  # column exists only in the Dukhan template
            raise TemplateError(f"Column '{header}' not found in the {plant_type} template.")
        targets.append((col, attr))

    lines = unique_lines(result.lines)  # each unique record once
    for i, line in enumerate(lines):
        r = BOM_FIRST_DATA_ROW + i
        for col, attr in targets:
            if attr.startswith("=LEN:"):
                src_col = idx[norm_header(attr[5:])]
                _put(ws, r, col, f"=LEN({get_column_letter(src_col)}{r})")
            else:
                value = getattr(line, attr)
                if value is not None:
                    _put(ws, r, col, value)

    # ---- sheet 2: SPIR TAG vs SUBMT
    ws2 = _sheet(wb, TAG_SHEET)
    _ensure_spir_type_column(ws2)
    idx2 = _header_index(ws2)
    for j, tag in enumerate(result.tags):
        r = TAG_FIRST_DATA_ROW + j
        for header, attr in zip(TAG_COLUMNS, TAG_ATTRS):
            if attr is None:
                continue
            value = getattr(tag, attr)
            if value is not None:
                _put(ws2, r, idx2[norm_header(header)], value)

    _border_table(ws, BOM_FIRST_DATA_ROW, len(lines))
    _border_table(ws2, TAG_FIRST_DATA_ROW, len(result.tags))
    _black_text(ws, BOM_FIRST_DATA_ROW, len(lines))
    _black_text(ws2, TAG_FIRST_DATA_ROW, len(result.tags))

    # freeze the header blocks: rows 1-5 on the BOM sheet, rows 1-2 on the tag sheet
    ws.freeze_panes = f"A{BOM_FIRST_DATA_ROW}"
    ws2.freeze_panes = f"A{TAG_FIRST_DATA_ROW}"

    # ---- sheet 3: VERIFICATION (lines to check by hand; the two template sheets are not affected)
    _write_verification(wb, ws, result.lines, lines)

    set_active_sheet(wb, ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- submission template
SUBMISSION_FIRST_DATA_ROW = 6  # rows 1-5 are the template header block on every sheet

# Columns not listed are left empty: External Number Assigned for Assembly, Safety Stock,
# Criticality, CDC_SAP CODED REMARK, CDC_CONSTRUCTION TYPE STATUS, CDC_BOM Parts Remark,
# MM requirement material type / group, PO text.
_COMMON = {
    "MATERIAL TEMP NUMBER": "temp_number",
    "MATERIAL TYPEVCATEGORY": "type_category",  # N for a spare with a 3/8 SAP number
    "OLD MATERIAL NUMBER/SPF NUMBER": "old_material_number",
    "MATERIAL DESCRIPTION": "description",
    "MAINTENANCE PLANNING PLANT": "planning_plant",
    "BASE UNIT OF MEASURE(T006)": "uom",
    "MANUFACTURERS PART NUMBER": "mfr_part_number",
    "MANUFACTURER NAME": "mfr_name",
    "CDC_SPIR NUMBER": "spir_number_full",  # with revision, as on the working template
    "CDC_ADD INFORMATION": "add_information_text",
}
SUBMISSION_SHEETS: dict[str, tuple[str, dict[str, str]]] = {
    # sheet -> (which lines, header -> BomLine attribute)
    "BOM HEADER": ("equipment", {
        **_COMMON,
        "MATERIAL IS AN EQUIPMENT INDICATOR": "equipment_indicator",
    }),
    "BOM PARTS": ("spares", {
        "BOM HEADER NUMBER": "bom_header_number",
        **_COMMON,
        "QUANTITY": "quantity",
        "POSITION NUMBER": "position_number",
        "SAP MATERIAL NUMBER": "sap_material_number",
    }),
    "MATERIAL MASTER": ("all", {
        **_COMMON,
        "MATERIAL IS AN EQUIPMENT INDICATOR": "equipment_indicator",
        "DELIVERY TIME IN WEEKS": "delivery_weeks",
        "MANUFACTURER COUNTRY CODE": "country_code",
        "MM REQUIREMENT MOVING AVERAGE PRICE": "price",
        "REORDER POINT (MINIMUM QTY)": "min_qty",
        "MAXIMUM STOCK LEVEL(MAX QTY)": "max_qty",
        "SAP MATERIAL NUMBER": "sap_material_number",
    }),
}


def write_submission(result: BomResult, template_path: Path) -> bytes:
    """Fill the BOM submission template: BOM HEADER = equipment, BOM PARTS = spares (linked to their
    equipment by BOM header number), MATERIAL MASTER = every material."""
    wb = openpyxl.load_workbook(template_path)
    unique = unique_lines(result.lines)  # same records as the working template, each once
    groups = {
        "equipment": [ln for ln in unique if ln.equipment_indicator == "YES"],
        "spares": [ln for ln in unique if ln.equipment_indicator != "YES"],
        "all": unique,
    }
    for name, (group, columns) in SUBMISSION_SHEETS.items():
        ws = _sheet(wb, name)
        idx = _header_index(ws)
        targets = []
        for header, attr in columns.items():
            col = idx.get(norm_header(header))
            if col is None:
                raise TemplateError(f"Column '{header}' not found on sheet '{name}' of the submission template.")
            targets.append((col, attr))

        # clear anything left below the header block (the template can carry stray sample values)
        for row in ws.iter_rows(min_row=SUBMISSION_FIRST_DATA_ROW):
            for c in row:
                c.value = None

        # a sheet shows fewer columns, so records can still repeat there (e.g. one spare under two
        # equipment on MATERIAL MASTER) - each unique row of this sheet once
        lines = _unique_rows(groups[group], _value_attrs(columns))
        for i, line in enumerate(lines):
            r = SUBMISSION_FIRST_DATA_ROW + i
            for col, attr in targets:
                value = getattr(line, attr)
                if value is not None:
                    _put(ws, r, col, value)
        _border_table(ws, SUBMISSION_FIRST_DATA_ROW, len(lines))
        _black_text(ws, SUBMISSION_FIRST_DATA_ROW, len(lines))
        ws.freeze_panes = f"A{SUBMISSION_FIRST_DATA_ROW}"

    set_active_sheet(wb, wb.worksheets[0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
