"""Combine several stored conversion outputs into one workbook ("Combine" on the History page).

The files are the exact working / submission templates the conversions produced, so nothing is
re-converted (no temp numbers are allocated). The first file is the base; the data rows of every
data sheet of the following files are appended below its rows, column by column header. Formulas
(the LEN() checks) are moved to their new row, the serial number is renumbered 1..n, and the data
table gets the same borders and frozen header block as a single conversion.
"""
from __future__ import annotations

import io
from collections.abc import Callable
from copy import copy

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.worksheet import Worksheet

from .excel_writer import (BOM_FIRST_DATA_ROW, BOM_SHEET, SUBMISSION_FIRST_DATA_ROW, SUBMISSION_SHEETS,
                           TAG_FIRST_DATA_ROW, TAG_SHEET, _black_text, _border_table, _header_index, _sheet,
                           set_active_sheet)
from .input_reader import norm_header

WORKING = "WORKING"
SUBMISSION = "SUBMISSION"

# kind -> [(sheet name, first data row, header of the column renumbered 1..n)]
SHEETS: dict[str, list[tuple[str, int, str | None]]] = {
    WORKING: [(BOM_SHEET, BOM_FIRST_DATA_ROW, "SERIAL NUMBER"), (TAG_SHEET, TAG_FIRST_DATA_ROW, None)],
    SUBMISSION: [(name, SUBMISSION_FIRST_DATA_ROW, None) for name in SUBMISSION_SHEETS],
}


def _data_rows(ws: Worksheet, first_row: int) -> list[int]:
    """Rows below the header block that hold at least one value (pre-formatted empty rows are skipped)."""
    return [row[0].row for row in ws.iter_rows(min_row=first_row, max_row=ws.max_row)
            if any(c.value is not None for c in row)]


def _copy_cell(src, dst) -> None:
    value = src.value
    if isinstance(value, str) and value.startswith("="):
        value = Translator(value, origin=src.coordinate).translate_formula(dst.coordinate)
    dst.value = value
    if src.has_style:  # style objects are copied one by one - style ids differ between workbooks
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.number_format = copy(src.alignment), src.number_format


def _append(base: Worksheet, src: Worksheet, first_row: int, next_row: int) -> int:
    """Append the data rows of src below next_row - 1 in base; returns the next free row."""
    base_cols = _header_index(base)
    col_map = {c: base_cols[h] for h, c in _header_index(src).items() if h in base_cols}
    for r in _data_rows(src, first_row):
        for src_col, dst_col in col_map.items():
            cell = src.cell(r, src_col)
            if cell.value is not None:
                _copy_cell(cell, base.cell(next_row, dst_col))
        next_row += 1
    return next_row


def combine(files: list[bytes], kind: str, on_progress: Callable[[int], None] | None = None) -> bytes:
    """files: the stored workbooks in the order they should appear. on_progress(n) after each file."""
    if kind not in SHEETS:
        raise ValueError(f"Unknown template kind {kind!r}")
    base = openpyxl.load_workbook(io.BytesIO(files[0]))
    names = [name for name, _, _ in SHEETS[kind]]
    sheets = [(_sheet(base, name), first, serial) for name, first, serial in SHEETS[kind]]
    next_rows = [(_data_rows(ws, first) or [first - 1])[-1] + 1 for ws, first, _ in sheets]
    if on_progress:
        on_progress(1)

    for done, content in enumerate(files[1:], start=2):
        wb = openpyxl.load_workbook(io.BytesIO(content))
        for i, (ws, first, _) in enumerate(sheets):
            next_rows[i] = _append(ws, _sheet(wb, names[i]), first, next_rows[i])
        wb.close()
        if on_progress:
            on_progress(done)

    for (ws, first, serial), next_row in zip(sheets, next_rows):
        count = next_row - first
        if serial:
            col = _header_index(ws).get(norm_header(serial))
            if col:
                for n in range(count):
                    ws.cell(first + n, col, n + 1)
        _border_table(ws, first, count)
        _black_text(ws, first, count)  # files made before this rule carry the template's red LEN() cells
        ws.freeze_panes = f"A{first}"

    # files made before the single-tab fix open grouped (Filter disabled) - keep their active sheet, alone
    set_active_sheet(base, base.active)
    buf = io.BytesIO()
    base.save(buf)
    return buf.getvalue()
