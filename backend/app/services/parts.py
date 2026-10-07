"""Parts Master: searchable copy of every extracted BOM line, plus search across the material master.

* save_lines()  - called by the converter for every successful conversion
* backfill()    - once at start-up: conversions made before this table existed are read back from
                  their stored working template, so older history is searchable too
* search()      - one field, a value with * wildcards; searches the conversion lines AND the
                  material_master_import rows. A part number found in both is merged into one row.
"""
from __future__ import annotations

import io
import logging

import openpyxl
from sqlalchemy import String, case, cast, func, literal, null, select, true
from sqlalchemy.orm import Session, undefer

from ..models import ConversionJob, ExtractedLine, material_master_import as master
from .bom_builder import BomResult, part_number_from_add_information
from .excel_writer import (BOM_COLUMNS, BOM_FIRST_DATA_ROW, BOM_SHEET, TAG_FIRST_DATA_ROW, TAG_SHEET,
                           _header_index, _sheet)
from .input_reader import norm_header

log = logging.getLogger(__name__)


def _str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


# ------------------------------------------------------------------ store
def save_lines(db: Session, job_id: int, result: BomResult) -> None:
    db.add_all(ExtractedLine(
        job_id=job_id, serial=line.serial, category=line.category, tag=_str(line.tag),
        sap_material_number=_str(line.sap_material_number), temp_number=_str(line.temp_number),
        old_material_number=_str(line.old_material_number), description=_str(line.description),
        planning_plant=_str(line.planning_plant), uom=_str(line.uom),
        # full part number (the template cuts it to 35 characters and keeps the rest in CDC_ADD INFORMATION)
        mfr_part_number=_str(line.add_information or line.mfr_part_number), mfr_name=_str(line.mfr_name),
        country_name=_str(line.country_name), spir_number=_str(line.spir_number), quantity=_str(line.quantity),
    ) for line in result.lines)


def _lines_from_workbook(job_id: int, content: bytes) -> list[ExtractedLine]:
    """Read the BOM lines back from a stored working template (for conversions made before this table)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = _sheet(wb, BOM_SHEET)
    cols = {attr: col for header, col in _header_index(ws).items()
            for h, attr in BOM_COLUMNS.items() if norm_header(h) == header and not attr.startswith("=")}
    tag_ws = _sheet(wb, TAG_SHEET)
    tcols = _header_index(tag_ws)
    tags = {}
    if norm_header("EQFNR") in tcols and norm_header("SUBMT") in tcols:
        for r in range(TAG_FIRST_DATA_ROW, tag_ws.max_row + 1):
            submt = _str(tag_ws.cell(r, tcols[norm_header("SUBMT")]).value)
            if submt:
                tags[submt] = _str(tag_ws.cell(r, tcols[norm_header("EQFNR")]).value)

    def get(r: int, attr: str):
        return _str(ws.cell(r, cols[attr]).value) if attr in cols else None

    lines = []
    for r in range(BOM_FIRST_DATA_ROW, ws.max_row + 1):
        if not (get(r, "temp_number") or get(r, "description") or get(r, "mfr_part_number")):
            continue
        lines.append(ExtractedLine(
            job_id=job_id, serial=len(lines) + 1, category=get(r, "category"),
            tag=tags.get(get(r, "bom_header_number") or ""), sap_material_number=get(r, "sap_material_number"),
            temp_number=get(r, "temp_number"), old_material_number=get(r, "old_material_number"),
            description=get(r, "description"), planning_plant=get(r, "planning_plant"), uom=get(r, "uom"),
            mfr_part_number=part_number_from_add_information(get(r, "add_information_text"))
            or get(r, "mfr_part_number"), mfr_name=get(r, "mfr_name"),
            country_name=get(r, "country_name"), spir_number=get(r, "spir_number"), quantity=get(r, "quantity"),
        ))
    return lines


def backfill(db: Session) -> int:
    """Add the lines of successful conversions that have none yet; returns how many conversions were filled."""
    has_lines = select(ExtractedLine.job_id).where(ExtractedLine.job_id == ConversionJob.id).exists()
    jobs = db.execute(select(ConversionJob).where(ConversionJob.status == "SUCCESS", ~has_lines)
                      .options(undefer(ConversionJob.output_file))).scalars().all()
    filled = 0
    for job in jobs:
        if not job.output_file:
            continue
        try:
            db.add_all(_lines_from_workbook(job.id, job.output_file))
            db.commit()
            filled += 1
        except Exception:  # one unreadable old file must not stop the others (or the start-up)
            db.rollback()
            log.exception("Parts Master: could not read the lines of conversion %s", job.id)
    return filled


# ------------------------------------------------------------------ search
L = ExtractedLine
# field -> (column on the conversion lines, column on material_master_import or None when it has no such column)
FIELDS = {
    "part_number": (L.mfr_part_number, master.c.manufacturers_part_number),
    "tag_number": (L.tag, None),
    "sap_material_number": (L.sap_material_number, master.c.sap_material_number),
    "material_temp_number": (L.temp_number, cast(master.c.material_temp_number, String)),
    "material_category": (L.category, master.c.material_type_category),
    "description": (L.description, master.c.material_description),
    "planning_plant": (L.planning_plant, master.c.maintenance_planning_plant),
    "manufacturer_name": (L.mfr_name, master.c.manufacturer_name),
    "country_name": (L.country_name, None),
    "spir": (L.spir_number, master.c.cdc_spir_number),
}


def like_pattern(value: str) -> str:
    """'12*' -> '12%' with LIKE's own % _ \\ escaped, so only * is a wildcard. No * = whole value."""
    escaped = value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
    return escaped.replace("*", "%")


# columns of a result row, in the order the merged values are combined
RESULT_FIELDS = ["material_temp_number", "material_category", "part_number", "description", "tag_number",
                 "sap_material_number", "planning_plant", "manufacturer_name", "country_name", "spir"]
# extra columns only the Excel download uses (Material Master sheet of the submission template)
EXPORT_FIELDS = ["old_material_number", "uom", "equipment_indicator", "min_qty", "max_qty"]


def part_key(value: str | None) -> str | None:
    """Same part number = same key: case and extra spaces are ignored (as in the numbering)."""
    return " ".join(value.split()).upper() if value and value.strip() else None


def _sql_part_key(column):
    return func.upper(func.regexp_replace(func.trim(column), r"\s+", " ", "g"))


def _conversion_rows(db: Session, condition) -> list[dict]:
    q = (
        select(literal("CONVERSION").label("source"), L.job_id.label("job_id"),
               L.temp_number.label("material_temp_number"), L.category.label("material_category"),
               L.mfr_part_number.label("part_number"), L.description.label("description"),
               L.tag.label("tag_number"), L.sap_material_number.label("sap_material_number"),
               L.planning_plant.label("planning_plant"), L.mfr_name.label("manufacturer_name"),
               L.country_name.label("country_name"), L.spir_number.label("spir"),
               ConversionJob.created_at.label("created_at"),
               L.old_material_number.label("old_material_number"), L.uom.label("uom"),
               case((L.category == "B", "YES"), (L.category == "L", "NO")).label("equipment_indicator"),
               null().label("min_qty"), null().label("max_qty"))
        .join(ConversionJob, ConversionJob.id == L.job_id)
        .where(condition)
        .order_by(ConversionJob.created_at.desc(), L.job_id.desc(), L.serial)
    )
    return [dict(r) for r in db.execute(q).mappings()]


def _master_rows(db: Session, condition) -> list[dict]:
    q = (
        select(literal("MASTER").label("source"), null().label("job_id"),
               cast(master.c.material_temp_number, String).label("material_temp_number"),
               master.c.material_type_category.label("material_category"),
               master.c.manufacturers_part_number.label("part_number"),
               master.c.material_description.label("description"), null().label("tag_number"),
               master.c.sap_material_number.label("sap_material_number"),
               master.c.maintenance_planning_plant.label("planning_plant"),
               master.c.manufacturer_name.label("manufacturer_name"), null().label("country_name"),
               master.c.cdc_spir_number.label("spir"), null().label("created_at"),
               master.c.old_material_number_spf_number.label("old_material_number"),
               master.c.base_unit_of_measure.label("uom"),
               case((master.c.material_is_equipment_indicator.is_(True), "YES"),
                    (master.c.material_is_equipment_indicator.is_(False), "NO")).label("equipment_indicator"),
               cast(master.c.reorder_point_minimum_qty, String).label("min_qty"),
               cast(master.c.maximum_stock_level_max_qty, String).label("max_qty"))
        .where(condition)
        .order_by(master.c.material_temp_number)
    )
    return [dict(r) for r in db.execute(q).mappings()]


def merge_rows(rows: list[dict]) -> dict:
    """One row from the conversion and material-master records of the same part number.

    Per column the conversion values come first and the material master fills what is missing; a value
    found in several records is shown once, different values are all kept (", " separated)."""
    conv = [r for r in rows if r["source"] == "CONVERSION"]
    ordered = conv + [r for r in rows if r["source"] != "CONVERSION"]
    merged: dict = {"source": "BOTH", "job_id": conv[0]["job_id"] if conv else None,
                    "created_at": max((r["created_at"] for r in conv if r["created_at"]), default=None)}
    for f in RESULT_FIELDS + EXPORT_FIELDS:
        values: list[str] = []
        for r in ordered:
            v = r.get(f)
            v = v.strip() if isinstance(v, str) else v
            if v and v not in values:
                values.append(v)
        merged[f] = ", ".join(values) or None
    return merged


def search(db: Session, field: str, value: str, limit: int, offset: int) -> tuple[list[dict], int]:
    """One page of find_parts()."""
    out = find_parts(db, field, value)
    return out[offset:offset + limit], len(out)


def find_parts(db: Session, field: str | None = None, value: str | None = None) -> list[dict]:
    """All matching conversion lines and material-master rows (every record when no value is given).
    A part number found in both sources is merged into ONE row (source BOTH) - its counterpart is looked up
    even when only one side matched the search, so e.g. a tag search still shows the SAP number that only
    the material master has."""
    if field and value and value.strip():
        line_col, master_col = FIELDS[field]
        pattern = like_pattern(value.strip())
        rows = _conversion_rows(db, line_col.ilike(pattern, escape="\\"))
        if master_col is not None:
            rows += _master_rows(db, master_col.ilike(pattern, escape="\\"))
    else:
        rows = _conversion_rows(db, true()) + _master_rows(db, true())

    keys = {"CONVERSION": set(), "MASTER": set()}
    for r in rows:
        if k := part_key(r["part_number"]):
            keys[r["source"]].add(k)
    # counterparts that did not match the search themselves
    if need := keys["CONVERSION"] - keys["MASTER"]:
        extra = _master_rows(db, _sql_part_key(master.c.manufacturers_part_number).in_(need))
        rows += extra
        keys["MASTER"] |= {part_key(r["part_number"]) for r in extra}
    if need := keys["MASTER"] - keys["CONVERSION"]:
        extra = _conversion_rows(db, _sql_part_key(L.mfr_part_number).in_(need))
        rows += extra
        keys["CONVERSION"] |= {part_key(r["part_number"]) for r in extra}
    both = keys["CONVERSION"] & keys["MASTER"]

    groups: dict[str, list[dict]] = {}
    for r in rows:
        k = part_key(r["part_number"])
        if k in both:
            groups.setdefault(k, []).append(r)

    # conversion results first (newest first), then material-master results; a merged row takes the place
    # of the first record of its part number
    out: list[dict] = []
    done: set[str] = set()
    for r in rows:
        k = part_key(r["part_number"])
        if k in both:
            if k not in done:
                done.add(k)
                out.append(merge_rows(groups[k]))
        else:
            out.append(r)
    return out
